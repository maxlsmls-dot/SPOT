# Returns attribution for the FTAI deck: re-runs the model for every on/off combination of the three theses.
# Usage: python3 returns_attribution.py <FTAI model .xlsx> <path to soffice.py wrapper>
import openpyxl, re, subprocess, sys, itertools, json, os
SRC = sys.argv[1]
SOFFICE = sys.argv[2]
LEVERS = {
    "margin": [49, 54, 61, 66],   # OEM escalator, pass-through, margin ex-PMA, PMA penetration
    "volume": [44, 32],           # FTAI share of MRO, SCI aircraft
    "power":  [75],               # units delivered (off = 0)
}
pat = re.compile(r'CHOOSE\(MATCH\(DCF!\$C\$3, \{"Bear","Base","Bull"\}, 0\),\s*([A-Z]+)(\d+),\s*([A-Z]+)(\d+),\s*([A-Z]+)(\d+)\)')
def build(on, path):
    wb = openpyxl.load_workbook(SRC)
    ws = wb["Assumptions"]
    for lever, rows in LEVERS.items():
        for r in rows:
            for c in ws[r]:
                v = c.value
                if isinstance(v, str) and "CHOOSE(MATCH(DCF!$C$3" in v:
                    m = pat.search(v)
                    assert m, (r, v)
                    col, bear, base = m.group(1), m.group(2), m.group(4)
                    if lever in on:
                        c.value = f"={col}{base}"
                    else:
                        c.value = 0 if lever == "power" else f"={col}{bear}"
    wb.save(path)
res = {}
os.makedirs("runs", exist_ok=True); os.makedirs("runs/out", exist_ok=True)
for k in range(4):
    for combo in itertools.combinations(LEVERS, k):
        name = "_".join(combo) or "none"
        p = f"runs/{name}.xlsx"
        build(set(combo), p)
        subprocess.run([sys.executable, SOFFICE, "--headless", "--convert-to", "xlsx", "--outdir", "runs/out", p], capture_output=True)
        wb = openpyxl.load_workbook(f"runs/out/{name}.xlsx", data_only=True)
        d = wb["DCF"]
        res[name] = dict(price=d["N29"].value, perp=d["N26"].value, exitm=d["N27"].value)
        print(name, res[name], flush=True)
json.dump(res, open("results.json", "w"), indent=1)
