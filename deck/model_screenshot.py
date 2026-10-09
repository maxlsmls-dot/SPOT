# Renders a cell range from the FTAI model to a cropped PNG for the appendix.
# Usage: python3 model_screenshot.py <model.xlsx> <soffice.py> <name> <sheet> <range> <out.png> [hide rows a-b] [hide cols C:F,...]
import openpyxl, subprocess, sys, os, glob
from PIL import Image, ImageChops
SRC, SOFFICE, name, sheet, rng, out = sys.argv[1:7]
hide = sys.argv[7] if len(sys.argv) > 7 else ""
hidecols = sys.argv[8] if len(sys.argv) > 8 else ""
wb = openpyxl.load_workbook(SRC)
for ws in wb.worksheets:
    if ws.title != sheet:
        ws.sheet_state = "hidden"
ws = wb[sheet]
wb.active = wb.worksheets.index(ws)
ws.print_area = rng
from openpyxl.utils import get_column_letter, range_boundaries
c1, r1, c2, r2 = range_boundaries(rng)
for c in range(c1, c2 + 1):
    L = get_column_letter(c)
    w = ws.column_dimensions[L].width or 9
    ws.column_dimensions[L].width = w * 1.22
for spec in filter(None, hidecols.split(",")):
    a, b = spec.split(":")
    for c in range(range_boundaries(a + "1")[0], range_boundaries(b + "1")[0] + 1):
        ws.column_dimensions[get_column_letter(c)].hidden = True
if hide and hide != "-":
    a, b = map(int, hide.split("-"))
    for r in range(a, b + 1):
        ws.row_dimensions[r].hidden = True
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_options.gridLines = False
ws.oddHeader.center.text = ""; ws.oddFooter.center.text = ""
os.makedirs("tmp", exist_ok=True)
x = f"tmp/{name}.xlsx"; wb.save(x)
subprocess.run([sys.executable, SOFFICE, "--headless", "--convert-to", "pdf", "--outdir", "tmp", x], capture_output=True)
pdf = f"tmp/{name}.pdf"
subprocess.run(["pdftoppm", "-png", "-r", "220", "-f", "1", "-l", "1", pdf, f"tmp/{name}"], check=True)
png = sorted(glob.glob(f"tmp/{name}-*.png"))[0]
im = Image.open(png).convert("RGB")
bg = Image.new("RGB", im.size, (255, 255, 255))
bbox = ImageChops.difference(im, bg).getbbox()
im = im.crop((max(bbox[0]-12,0), max(bbox[1]-12,0), min(bbox[2]+12, im.width), min(bbox[3]+12, im.height)))
im.save(out); print(out, im.size)
