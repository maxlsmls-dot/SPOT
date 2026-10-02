#!/usr/bin/env python3
"""Two audiobook tabs styled to match the SPOT operating model (RPM / Assumptions conventions):
   'Audiobook Inputs'     segmented inputs with FY columns, scenario rows for cost growth, Notes column
   'Audiobook GM Bridge'  B -> C -> A -> N -> contribution to gross margin and y/y headwind, FY'24-FY'27E
Output: data/SPOT_Audiobook_Tabs.xlsx.  Run: python3 analysis/audiobooks_model_tabs.py
"""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from audiobooks_model import INPUTS, compute, inject_cached_values, DATA
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

I = {k: v for k, _, v, *_ in INPUTS}
SH_IN, SH_BR = "Audiobook Inputs", "Audiobook GM Bridge"
IN = f"'{SH_IN}'!"
NAVY = PatternFill("solid", fgColor="FF002060"); BAND = PatternFill("solid", fgColor="FFE7E6E6")
GOLD = PatternFill("solid", fgColor="FFFFF2CC"); YELLOW = PatternFill("solid", fgColor="FFFFFF00")
F_HDR = Font(name="Calibri", size=11, bold=True, color="FFFFFFFF")
F_TXT = Font(name="Calibri", size=11); F_B = Font(name="Calibri", size=11, bold=True); F_I = Font(name="Calibri", size=11, italic=True)
F_IN = Font(name="Calibri", size=11, color="FF0070C0"); F_IN_B = Font(name="Calibri", size=11, bold=True, color="FF0070C0")
THIN = Side(style="thin"); CENTER = Alignment(horizontal="center"); LEFT = Alignment(horizontal="left")
NOTE = Alignment(wrap_text=True, vertical="top")
FM = '#,##0_);\\(#,##0\\)'; FBP = '0_);\\(0\\)'; FP0 = '0%'; FP1 = '0.0%'; F0 = '0'; FX = '0.00'; FRATE = '0.00000'
YC = ["C", "D", "E", "F"]; YEARS = ["FY'24", "FY'25", "FY'26E", "FY'27E"]

def setup(ws, title):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 8.9; ws.column_dimensions["B"].width = 58
    for c in YC + ["G"]: ws.column_dimensions[c].width = 13
    ws.column_dimensions["H"].width = 78
    ws["B2"] = title
    for col in ["B"] + YC + ["G", "H"]:
        ws[f"{col}2"].fill = NAVY; ws[f"{col}2"].font = F_HDR
    for col, y in zip(YC, YEARS): ws[f"{col}2"] = y; ws[f"{col}2"].alignment = CENTER
    ws["H2"] = "Notes"
    for col in ["B"] + YC + ["G", "H"]: ws[f"{col}3"].fill = BAND
    ws.freeze_panes = "A4"

class Writer:
    def __init__(self, ws): self.ws = ws; self.r = 3
    def row(self, label=None, vals=None, note=None, bold=False, italic=False, fill=None, top=False, box=False, g=None, fmt=None, align=CENTER, label_font=None):
        self.r += 1; ws = self.ws; rr = self.r
        if label is not None:
            c = ws.cell(row=rr, column=2, value=label); c.font = label_font or (F_B if bold else (F_I if italic else F_TXT))
            if italic: c.alignment = LEFT
        if vals:
            for col, v in zip(YC, vals):
                if v is None or v == "": continue
                self._cell(f"{col}{rr}", v, bold, fmt, align)
        if g is not None: self._cell(f"G{rr}", g, bold, fmt, align)
        if note:
            c = ws.cell(row=rr, column=8, value=note); c.font = F_I; c.alignment = NOTE
        if fill:
            for col in ["B"] + YC: ws[f"{col}{rr}"].fill = fill
        if top:
            for col in ["B"] + YC: ws[f"{col}{rr}"].border = Border(top=THIN)
        if box:
            ws[f"B{rr}"].border = Border(left=THIN, top=THIN, bottom=THIN)
            for col in YC: ws[f"{col}{rr}"].border = Border(top=THIN, bottom=THIN)
        return rr
    def _cell(self, ref, v, bold, fmt, align):
        c = self.ws[ref]; c.value = v; c.alignment = align
        is_formula = isinstance(v, str) and v.startswith("=")
        c.font = (F_B if bold else F_TXT) if is_formula else (F_IN_B if bold else F_IN)
        if fmt: c.number_format = fmt
    def blank(self): self.r += 1
    def section(self, label): return self.row(label, bold=True)
    def subhdr(self, labels, g=None):
        rr = self.row(labels[0], label_font=F_I)
        for col, t in zip(YC, labels[1:]):
            c = self.ws[f"{col}{rr}"]; c.value = t; c.font = F_B; c.alignment = CENTER; c.border = Border(bottom=THIN)
        if g: c = self.ws[f"G{rr}"]; c.value = g; c.font = F_B; c.alignment = CENTER; c.border = Border(bottom=THIN)
        return rr

def build():
    r = compute()
    wb = Workbook(); wb.properties.creator = "SPOT"; wb.properties.lastModifiedBy = "SPOT"
    # ======================================================================= Inputs
    ws = wb.active; ws.title = SH_IN; setup(ws, "Audiobook Assumptions"); w = Writer(ws); A = {}
    def inp(key, label, vals, note, fmt, placeholder_cols=(), bold=False):
        rr = w.row(label, vals, note, fmt=fmt, bold=bold)
        for col in placeholder_cols: ws[f"{col}{rr}"].fill = GOLD
        return rr
    rr = w.row("Case selector (repoint to Valuation!$C$3 when the tabs are added to the model)", ["Base"], "Drives the licensing-cost growth assumption in section E. Bear = adverse for SPOT (faster cost growth).", fmt=None)
    ws[f"C{rr}"].fill = GOLD; A["case"] = f"{IN}$C${rr}"
    w.blank()
    w.section("A. Group revenue, €M (denominator for basis points)")
    rr = inp("grp", "Total revenue", [I["grp_rev_2024"], I["grp_rev_2025"], I["grp_rev_2026"], I["grp_rev_2027"]],
             "FY'24-FY'25: Form 20-F. FY'26E-FY'27E: placeholders, link to IS total revenue.", FM, placeholder_cols=["E", "F"])
    for col, k in zip(YC, ["grp_rev_2024", "grp_rev_2025", "grp_rev_2026", "grp_rev_2027"]): A[k] = f"{IN}${col}${rr}"
    w.blank()
    w.section("B. Bundle royalty saving: MLC contingency disclosed in Spotify filings, €M cumulative from 1 Mar 2024")
    series = [("mlc_cum_q224", "6-K Q2-24: cumulative to 30 Jun 2024", "6-K Q2-24 legal proceedings note: 'approximately €46 million, of which approximately €35 million relates to the three months ended June 30, 2024'. Additional royalties due to the MLC if Premium were not a bundle."),
              ("mlc_q224_only", "   of which Apr-Jun 2024", "Same sentence; implies March 2024 alone was about €11M."),
              ("mlc_cum_q324", "6-K Q3-24: cumulative to 30 Sep 2024", "6-K Q3-24 legal proceedings note."),
              ("mlc_cum_q424", "20-F FY24: cumulative to 31 Dec 2024", "Form 20-F FY2024, legal proceedings."),
              ("mlc_cum_q125", "6-K Q1-25: cumulative to 31 Mar 2025", "6-K Q1-25 legal proceedings note."),
              ("mlc_cum_q225", "6-K Q2-25: cumulative to 30 Jun 2025", "6-K Q2-25 legal proceedings note."),
              ("mlc_cum_q325", "6-K Q3-25: cumulative to 30 Sep 2025", "6-K Q3-25: first filing to add 'any liability would be partially offset by direct deals with publishers'."),
              ("mlc_cum_q425", "20-F FY25: cumulative to 31 Dec 2025", "Form 20-F FY2025, legal proceedings."),
              ("mlc_cum_q126", "6-K Q1-26: cumulative to 31 Mar 2026", "6-K Q1-26 (about $473M at the time; NMPA cited 'nearly $480 million by Spotify's own admission')."),
              ("mlc_cum_q226", "6-K Q2-26: cumulative to 30 Jun 2026 (verify)", "Quoted by Music Ally, 7 Sep 2026, as €437M. A second reading of the same filing gives €473M, which is also the dollar value of the Q1-26 figure. Verify in the Q2-26 6-K before quoting.")]
    for k, label, note in series:
        rr = inp(k, label, [I[k]], note, FM); A[k] = f"{IN}$C${rr}"
    rr = inp("bundle_rr", "Gross bundle saving, forecast placeholders", ["", "", I["bundle_rr_2026"], I["bundle_rr_2027"]],
             "FY'26E: 4 x the average quarterly increment from Q2-25 to Q1-26 (€51M). FY'27E: flat; the Phonorecords IV settlement that permits the bundle rate runs to end-2027. Link to model.", FM, placeholder_cols=["E", "F"])
    A["bundle_rr_2026"] = f"{IN}$E${rr}"; A["bundle_rr_2027"] = f"{IN}$F${rr}"
    w.blank()
    w.section("C. Bundle saving cross-check: Billboard analysis of Spotify's MLC reports, paid tiers")
    for k, label, note, fmt in [("bb_q423_mech", "Mechanical royalties paid to the MLC, Q4-23, $M", "Billboard, 'How Much Have Spotify Bundles Decreased the Mechanical Per-Stream Rate? (Analysis)'.", '0.0'),
                                ("bb_q425_mech", "Mechanical royalties paid to the MLC, Q4-25, $M", "Same analysis; down 45%.", '0.0'),
                                ("bb_q423_rate", "Blended per-stream mechanical rate, Q4-23, $", "Same analysis.", FRATE),
                                ("bb_q425_rate", "Blended per-stream mechanical rate, Q4-25, $", "Same analysis; down 51%. Streams grew about 19bn between the two quarters.", FRATE),
                                ("eurusd_q425", "EUR/USD, Q4-25 average", "Approximate.", FX)]:
        rr = inp(k, label, [I[k]], note, fmt); A[k] = f"{IN}$C${rr}"
    w.blank()
    w.section("D. Direct publisher deals that hand part of the saving back (share of US mechanicals in FY'24 column; start month in FY'25 column)")
    for sk, st, label, note in [("share_umpg", "start_umpg", "Universal Music Publishing Group", "Announced 26 Jan 2025 (MBW); reported to nullify the bundle discount going forward. Share: estimate from US publishing market rankings (range 20-25%)."),
                                ("share_wcm", "start_wcm", "Warner Chappell Music", "Announced 6 Feb 2025 (MBW): 'supersedes the bundling payment structure'; terms 'continue to recognize a difference between bundled and music-only listeners'. Share: estimate (11-15%)."),
                                ("share_kobalt", "start_kobalt", "Kobalt", "Announced 13 Aug 2025 (CMU, MBW): 'bespoke terms rather than the compulsory licence rates'. Share: estimate (4-7%)."),
                                ("share_smp", "start_smp", "Sony Music Publishing", "Announced 18 Sep 2025 (MBW, Variety); SMP ranked #1 publisher in every quarter of 2025. Share: estimate (23-28%)."),
                                ("share_bmg", "start_bmg", "BMG", "Announced 9 Oct 2025 (CMU). Share: estimate (2-4%).")]:
        rr = w.row(label, [I[sk], I[st]], note); ws[f"C{rr}"].number_format = FP0; ws[f"D{rr}"].number_format = F0
        A[sk] = f"{IN}$C${rr}"; A[st] = f"{IN}$D${rr}"
    rr = inp("restoration", "Restoration of the bundle discount in the direct deals", [I["restoration"]],
             "Estimate. Spotify calls the offset 'partial'; the Warner terms still distinguish bundled from music-only listeners while payments are 'substantially improved'. Range 50-100%.", FP0); A["restoration"] = f"{IN}$C${rr}"
    rr = inp("mlc", "MLC amended complaint succeeds in FY'27 (1 = yes)", [I["mlc_loss_2027"]],
             "Scenario switch: the independents' share of the saving goes to zero. Dismissed Jan 2025; amended complaint Oct 2025; interlocutory appeal denied 1 Sep 2026; fact discovery cutoff 13 Mar 2027; no trial date.", F0); A["mlc_loss_2027"] = f"{IN}$C${rr}"
    w.blank()
    w.section("E. Audiobook licensing cost paid to book publishers (included and paid hours)")
    rr = inp("c_2024", "FY'24 cost anchor, €M", [I["c_2024"]],
             "Spotify via Axios, 15 Oct 2024: paying publishers 'hundreds of millions of dollars a year' (repeated Mar-25; 'tens of millions' by Feb-24 per Digital Music News). Read as about $205M for the full year at a 1.08 average rate. Cross-check: APA 2025 US publisher receipts $2.43B x 10-14% Spotify share x 1.4 non-US / 1.13 = €301-421M for 2025 vs €300M modelled.", FM); A["c_2024"] = f"{IN}$C${rr}"
    sel = A["case"]
    rr_g = w.row("Licensing cost growth, % y/y", ["", I["c_g_2025"], None, None],
                 "FY'25: Spotify newsroom Oct 2025, hours +37% like-for-like, times about 1.15 for the wider paying base (Family and Duo members, DACH, France and Benelux annualising). FY'26E: Investor Day May 2026, listeners +60%, cost at about 0.7x listener growth. FY'27E: assumed deceleration. Scenario rows below; Bear = faster cost growth.", fmt=FP1)
    bear = rr_g + 1; base = rr_g + 2; bull = rr_g + 3
    for col in ["E", "F"]:
        c = ws[f"{col}{rr_g}"]; c.value = f'=IF({sel}="Bear",{col}{bear},IF({sel}="Bull",{col}{bull},{col}{base}))'; c.font = F_TXT; c.alignment = CENTER; c.number_format = FP1; c.fill = GOLD
    for name, vals, bd in [("Bear", [0.60, 0.45], Border(left=THIN, top=THIN)), ("Base", [0.43, 0.30], Border(left=THIN)), ("Bull", [0.30, 0.20], Border(left=THIN, bottom=THIN))]:
        rr = w.row(name, ["", "", vals[0], vals[1]], fmt=FP1); ws[f"B{rr}"].border = bd
    A["c_g_2025"] = f"{IN}$D${rr_g}"; A["c_g_2026"] = f"{IN}$E${rr_g}"; A["c_g_2027"] = f"{IN}$F${rr_g}"
    w.blank()
    w.section("F. Paid audiobook revenue, €M (Audiobooks+, top-ups, a la carte)")
    rr = inp("a_rev", "Paid audiobook revenue", [I["a_rev_2024"], I["a_rev_2025"], I["a_rev_2026"], I["a_rev_2027"]],
             "FY'24: top-ups only from Mar-24, estimate. FY'25: US Audiobooks+ from Aug-25 at $11.99 plus 11 EU markets late in the year, estimate. FY'26E: $100M ARR reached Jul-26 (Q2-26 call) is about €88M run-rate; full year lower. FY'27E: +60%. Link to model.", FM, placeholder_cols=["E", "F"])
    for col, k in zip(YC, ["a_rev_2024", "a_rev_2025", "a_rev_2026", "a_rev_2027"]): A[k] = f"{IN}${col}${rr}"
    w.row("   Context: Audiobooks+ ARR, $M (Jul-26) and paying users, M (May-26)", [I["ab_plus_arr"], I["ab_plus_payers"]], "Investor Day 21 May 2026 via Publishers Weekly: 'more than one million' payers, 'on track for $100M ARR in July'; Q2-26 call: reached.", fmt='0.0')
    w.blank()
    w.section("G. Consensus gross-margin expansion, bp (context only)")
    rr = inp("cons", "Consensus y/y gross-margin expansion", ["", "", I["cons_gm_exp_2026"], I["cons_gm_exp_2027"]], "Placeholders: 33.2% vs 32.0% and ~34.5% vs 33.2%. Link to the consensus tab.", FBP, placeholder_cols=["E", "F"])
    A["cons_gm_exp_2026"] = f"{IN}$E${rr}"; A["cons_gm_exp_2027"] = f"{IN}$F${rr}"
    w.blank()
    w.section("H. Memo, excluded from the bridge")
    rr = inp("p_2024", "June 2024 US price increase, annualised revenue, €M", [I["p_2024"]], "Estimate (~7% blended on US Premium revenue). The item management credits to audiobooks; a pricing effect, so it stays outside N.", FM); A["p_2024"] = f"{IN}$C${rr}"
    w.blank()
    w.row("Colour key: blue = hard-coded input (notes give source or derivation); light gold = forecast assumption or placeholder to link to the model; black = formula.", label_font=F_I)

    # ======================================================================= Bridge
    bs = wb.create_sheet(SH_BR); setup(bs, "Audiobook Gross Margin Bridge"); b = Writer(bs)
    b.row("€ in millions unless stated.  N = B + A - C:  B = bundle royalty saving retained,  A = paid audiobook revenue,  C = audiobook licensing cost paid to book publishers.", label_font=F_I)
    b.blank()
    b.section("1. Bundle Royalty Saving (B)")
    b.subhdr(["MLC contingency by filing, €M", "Cumulative", "Increment", "Months", "€M / month"])
    keys = [("6-K Q2-24 (30 Jun 2024)", "mlc_cum_q224", 4), ("6-K Q3-24 (30 Sep 2024)", "mlc_cum_q324", 3), ("20-F FY24 (31 Dec 2024)", "mlc_cum_q424", 3),
            ("6-K Q1-25 (31 Mar 2025)", "mlc_cum_q125", 3), ("6-K Q2-25 (30 Jun 2025)", "mlc_cum_q225", 3), ("6-K Q3-25 (30 Sep 2025)", "mlc_cum_q325", 3),
            ("20-F FY25 (31 Dec 2025)", "mlc_cum_q425", 3), ("6-K Q1-26 (31 Mar 2026)", "mlc_cum_q126", 3), ("6-K Q2-26 (30 Jun 2026): verify", "mlc_cum_q226", 3)]
    R = {}; prev = None
    for label, k, months in keys:
        rr = b.r + 1; inc = f"=C{rr}" if prev is None else f"=C{rr}-C{prev}"
        b.row(label, [f"={A[k]}", inc, months, f"=D{rr}/E{rr}"], fmt=FM); bs[f"E{rr}"].number_format = F0; bs[f"F{rr}"].number_format = '0.0'; bs[f"E{rr}"].font = F_IN
        R[k] = rr; prev = rr
    rr = b.r + 1; b.row("   of which Apr-Jun 2024 (March 2024 alone about €11M)", [f"={A['mlc_q224_only']}", "", 3, f"=C{rr}/E{rr}"], fmt=FM); bs[f"E{rr}"].number_format = F0; bs[f"F{rr}"].number_format = '0.0'; bs[f"E{rr}"].font = F_IN
    R["avg"] = b.row("Average quarterly increment, Q2-25 to Q1-26 (x4 annualised in next column)", [f"=AVERAGE(D{R['mlc_cum_q225']}:D{R['mlc_cum_q126']})", f"=C{b.r+1}*4"], "Increments ran €50-56M a quarter with no step down after the direct deals, so the figure is gross: every paid-tier stream is still reported through the MLC at the bundle rate and the direct-licensed publishers receive a top-up outside it.", fmt='0.0'); bs[f"D{R['avg']}"].number_format = FM
    b.blank()
    b.subhdr(["Billboard cross-check (paid tiers)", "Q4-23", "Q4-25", "", ""])
    rb1 = b.row("Mechanical royalties paid to the MLC, $M", [f"={A['bb_q423_mech']}", f"={A['bb_q425_mech']}"], fmt='0.0')
    rb2 = b.row("Blended per-stream rate, $", [f"={A['bb_q423_rate']}", f"={A['bb_q425_rate']}"], fmt=FRATE)
    rb3 = b.row("Implied streams, bn", [f"=C{rb1}/C{rb2}/1000", f"=D{rb1}/D{rb2}/1000"], fmt='0.0')
    rb4 = b.row("Implied Q4-25 discount: streams x Q4-23 rate less paid, $M (col D); in €M (col E) vs 20-F FY25 increment (col F)",
                ["", f"=D{rb3}*C{rb2}*1000-D{rb1}", f"=D{b.r+1}/{A['eurusd_q425']}", f"=D{R['mlc_cum_q425']}"], "Two sources that share no inputs agree on a gross discount of about €50M a quarter. Streams grew between the quarters, so direct-licensed publishers have not left the blanket licence.", fmt='0.0')
    b.blank()
    b.subhdr(["Gross saving, hand-back, B", "FY'24", "FY'25", "FY'26E", "FY'27E"])
    R["gross"] = b.row("Gross bundle saving", [f"={A['mlc_cum_q424']}", f"={A['mlc_cum_q425']}-{A['mlc_cum_q424']}", f"={A['bundle_rr_2026']}", f"={A['bundle_rr_2027']}"], "FY'24 and FY'25 straight from the filings; FY'26E-FY'27E are placeholders on the inputs tab.", bold=True, top=True, fmt=FM)
    b.subhdr(["Months on direct licence by publisher", "", "", "", ""], g="Share")
    pubs = {}
    for label, sk, st in [("Universal Music Publishing Group", "share_umpg", "start_umpg"), ("Warner Chappell Music", "share_wcm", "start_wcm"), ("Kobalt", "share_kobalt", "start_kobalt"), ("Sony Music Publishing", "share_smp", "start_smp"), ("BMG", "share_bmg", "start_bmg")]:
        rr = b.row("   " + label, [0, f"=13-{A[st]}", 12, 12], g=f"={A[sk]}", fmt=F0); bs[f"G{rr}"].number_format = FP0; bs[f"C{rr}"].font = F_IN; pubs[sk] = rr
    pf, pl = min(pubs.values()), max(pubs.values())
    R["S"] = b.row("   Sum of shares on direct licences (independents are the remainder)", g=f"=SUM(G{pf}:G{pl})", fmt=FP0)
    R["rest"] = b.row("   Restoration of the discount in the direct deals", g=f"={A['restoration']}", fmt=FP0)
    R["hb"] = b.row("Hand-back = gross x sum(share x months / 12) x restoration", [f"={c}{R['gross']}*SUMPRODUCT($G${pf}:$G${pl},{c}{pf}:{c}{pl})/12*$G${R['rest']}" for c in YC], fmt=FM)
    R["B"] = b.row("B: Net bundle saving retained by Spotify", [f"={c}{R['gross']}-{c}{R['hb']}" for c in YC[:3]] + [f"=IF({A['mlc_loss_2027']}=1,F{R['gross']}*$G${R['S']}*(1-$G${R['rest']}),F{R['gross']}-F{R['hb']})"],
                   "FY'27E: if the MLC switch on the inputs tab is 1, the independents' share of the saving goes to zero.", bold=True, top=True, fmt=FM)
    b.blank()
    b.section("2. Audiobook Licensing Cost (C)")
    R["g"] = b.row("% y/y Growth", ["", f"={A['c_g_2025']}", f"={A['c_g_2026']}", f"={A['c_g_2027']}"], "FY'26E-FY'27E follow the case selector on the inputs tab (Bear / Base / Bull rows).", italic=True, fmt=FP1)
    rc = b.r + 1
    R["C"] = b.row("C: Licensing cost paid to book publishers", [f"={A['c_2024']}", f"=C{rc}*(1+D{R['g']})", f"=D{rc}*(1+E{R['g']})", f"=E{rc}*(1+F{R['g']})"], "FY'24 anchored on Spotify's 'hundreds of millions of dollars a year' (Oct-24); grown by consumption.", bold=True, top=True, fmt=FM)
    b.blank()
    b.section("3. Paid Audiobook Revenue (A)")
    R["A"] = b.row("A: Paid audiobook revenue", [f"={A[k]}" for k in ["a_rev_2024", "a_rev_2025", "a_rev_2026", "a_rev_2027"]], "Counted in full: C is grown on total consumption, so it already carries the publisher cost of paid hours.", bold=True, top=True, fmt=FM)
    b.blank()
    b.section("4. Net Contribution to Gross Profit and Gross Margin")
    R["B2"] = b.row("B: Net bundle saving retained", [f"={c}{R['B']}" for c in YC], fmt=FM)
    R["A2"] = b.row("A: Paid audiobook revenue", [f"={c}{R['A']}" for c in YC], fmt=FM)
    R["C2"] = b.row("Less C: Licensing cost paid to book publishers", [f"=-{c}{R['C']}" for c in YC], fmt=FM)
    R["N"] = b.row("Audiobook Contribution to Gross Profit (N)", [f"={c}{R['B2']}+{c}{R['A2']}+{c}{R['C2']}" for c in YC], "Add this row to IS gross profit (FY columns) to embed the audiobook effect at the level.", bold=True, fill=YELLOW, box=True, fmt=FM)
    R["rev"] = b.row("Total Revenue", [f"={A[k]}" for k in ["grp_rev_2024", "grp_rev_2025", "grp_rev_2026", "grp_rev_2027"]], fmt=FM)
    R["Nbp"] = b.row("% Contribution to Gross Margin, bp", [f"={c}{R['N']}/{c}{R['rev']}*10000" for c in YC], "Level: what audiobooks did to gross margin in each year versus no audiobooks.", italic=True, fmt=FBP)
    R["dN"] = b.row("Y/y Change in Contribution, €M", ["", f"=D{R['N']}-C{R['N']}", f"=E{R['N']}-D{R['N']}", f"=F{R['N']}-E{R['N']}"], fmt=FM)
    R["dNbp"] = b.row("Y/y Gross Margin Headwind, bp", ["", f"=D{R['dN']}/D{R['rev']}*10000", f"=E{R['dN']}/E{R['rev']}*10000", f"=F{R['dN']}/F{R['rev']}*10000"],
                      "Subtract from the model's y/y gross-margin expansion (negative = drag). Use either this row or the €M contribution row, not both.", bold=True, fill=YELLOW, box=True, fmt=FBP)
    b.row("   of which music publishers reclaim the bundle saving (change in B)", ["", f"=D{R['B2']}-C{R['B2']}", f"=E{R['B2']}-D{R['B2']}", f"=F{R['B2']}-E{R['B2']}"], italic=True, fmt=FM)
    b.row("   of which book publishers paid on consumption (change in C)", ["", f"=D{R['C2']}-C{R['C2']}", f"=E{R['C2']}-D{R['C2']}", f"=F{R['C2']}-E{R['C2']}"], italic=True, fmt=FM)
    b.row("   of which paid audiobook revenue (change in A)", ["", f"=D{R['A2']}-C{R['A2']}", f"=E{R['A2']}-D{R['A2']}", f"=F{R['A2']}-E{R['A2']}"], italic=True, fmt=FM)
    R["cons"] = b.row("Consensus y/y gross-margin expansion, bp (context)", ["", "", f"={A['cons_gm_exp_2026']}", f"={A['cons_gm_exp_2027']}"], fmt=FBP)
    b.row("Headwind as % of consensus expansion (context)", ["", "", f"=-E{R['dNbp']}/E{R['cons']}", f"=-F{R['dNbp']}/F{R['cons']}"], "Absolute drag, not delta to consensus: the Street already embeds some audiobook cost.", italic=True, fmt=FP0)
    b.row("Memo, excluded: June 2024 US price increase, annualised €M", [f"={A['p_2024']}"], fmt=FM)
    b.blank()
    b.section("5. Sensitivity: growth assigned to audiobook licensing cost (applied to FY'26E and FY'27E; B, A and revenue held)")
    b.subhdr(["Licensing cost growth p.a.", "N FY'26E, €M", "N FY'27E, €M", "Headwind FY'26E, bp", "Headwind FY'27E, bp"])
    C25 = f"$D${R['C']}"; N25 = f"$D${R['N']}"; B26, B27 = f"$E${R['B']}", f"$F${R['B']}"; A26, A27 = f"$E${R['A']}", f"$F${R['A']}"; V26, V27 = f"$E${R['rev']}", f"$F${R['rev']}"
    for g in [0.20, 0.30, 0.43, 0.50, 0.60]:
        rr = b.r + 1
        b.row(None, [g, f"={B26}+{A26}-{C25}*(1+C{rr})", f"={B27}+{A27}-{C25}*(1+C{rr})^2", f"=(D{rr}-{N25})/{V26}*10000"], g=f"=(E{rr}-D{rr})/{V27}*10000", fmt=FM)
        bs[f"C{rr}"].number_format = FP0; bs[f"F{rr}"].number_format = FBP; bs[f"G{rr}"].number_format = FBP
        bs.cell(row=rr, column=2, value="   Base growth, FY'26E" if abs(g - 0.43) < 1e-9 else None).font = F_I
    rr = b.row("   Bridge base case (43% in FY'26E, 30% in FY'27E)", ["", f"=E{R['N']}", f"=F{R['N']}", f"=E{R['dNbp']}"], "The headwind is negative in every row; its size scales with the cost growth, which rests on listeners +60% (disclosed) and a 0.7x haircut to cost (estimate).", g=f"=F{R['dNbp']}", bold=True, top=True, fmt=FM)
    bs[f"F{rr}"].number_format = FBP; bs[f"G{rr}"].number_format = FBP
    out = os.path.join(DATA, "SPOT_Audiobook_Tabs.xlsx"); wb.save(out); inject_cached_values(out); print("saved", out)
    return r

if __name__ == "__main__":
    build()
