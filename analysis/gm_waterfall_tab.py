#!/usr/bin/env python3
"""'GM Waterfall' tab in the user's own layout (labels in A, FY25-FY30 in B-G), FY'25 tied to reported COGS.

  Revenue (RPM row 58)  ->  Label % / Publisher %  ->  Royalties (normal, before Discovery Mode)
  Discovery Mode eligible streams x DM % of eligible  ->  DM royalties (label royalties on DM streams)  ->  DM netted royalties (after the 30% haircut)
  Audiobook licensing (Audiobook GM Bridge row 41)  +  other cost of revenue (plug in FY'25, held as % of revenue)
  COGS = royalties - DM royalties + DM netted royalties + audiobooks + other  ->  gross profit, gross margin

Output: data/SPOT_GM_Waterfall_Tab.xlsx with stub sheets (RPM, DM Inputs, Audiobook GM Bridge, IS) holding only the
referenced cells so it calculates stand-alone. After pasting into the model, Find & Replace "[SPOT_GM_Waterfall_Tab.xlsx]"
with nothing. Run: python3 analysis/gm_waterfall_tab.py  (then recalc with LibreOffice)
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "data", "SPOT_GM_Waterfall_Tab.xlsx")
F_TXT = Font(name="Calibri", size=11); F_B = Font(name="Calibri", size=11, bold=True); F_I = Font(name="Calibri", size=11, italic=True)
F_IN = Font(name="Calibri", size=11, color="FF0070C0")
GOLD = PatternFill("solid", fgColor="FFFFF2CC"); YELLOW = PatternFill("solid", fgColor="FFFFFF00")
THIN = Side(style="thin"); CENTER = Alignment(horizontal="center"); NOTE = Alignment(wrap_text=True, vertical="top")
FM = '#,##0_);\\(#,##0\\)'; FP1 = '0.0%'; FP0 = '0%'; FBP = '0_);\\(0\\)'
COLS = ["B", "C", "D", "E", "F", "G"]; YEARS = ["FY25", "FY26", "FY27", "FY28", "FY29", "FY30"]
RPM_COLS = ["Q", "V", "AA", "AB", "AC", "AD"]
RPM_TOTAL = [17186, 19410, 21172, 23166, 25017, 26482]
DMI = {"D30": 0.512, "E31": 0.52, "F31": 0.525, "D35": 0.33, "E36": 0.33, "F36": 0.335}
AB = {"D41": 300.2, "E41": 429.286, "F41": 558.072}

def build():
    wb = Workbook(); ws = wb.active; ws.title = "GM Waterfall"; ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 44
    for c in COLS: ws.column_dimensions[c].width = 11
    ws.column_dimensions["H"].width = 80
    r = [1]
    def row(label, vals=None, fmt=None, bold=False, italic=False, note=None, top=False, fill=None):
        rr = r[0]
        if label is not None:
            c = ws[f"A{rr}"]; c.value = label; c.font = F_B if bold else (F_I if italic else F_TXT)
        if vals:
            for col, v in zip(COLS, vals):
                if v is None or v == "": continue
                c = ws[f"{col}{rr}"]; c.value = v; c.alignment = CENTER
                c.font = (F_B if bold else F_TXT) if (isinstance(v, str) and v.startswith("=")) else F_IN
                if fmt: c.number_format = fmt
        if note: n = ws[f"H{rr}"]; n.value = note; n.font = F_I; n.alignment = NOTE
        if top:
            for col in ["A"] + COLS: ws[f"{col}{rr}"].border = Border(top=THIN)
        if fill:
            for col in COLS: ws[f"{col}{rr}"].fill = fill
        r[0] += 1; return rr
    def blank(): r[0] += 1
    def flat(first, rr):   # FY26-FY30 = prior year, gold
        for i, col in enumerate(COLS[1:], 1):
            c = ws[f"{col}{rr}"]; c.value = f"={COLS[i-1]}{rr}"; c.font = F_TXT; c.alignment = CENTER; c.number_format = ws[f"B{rr}"].number_format; c.fill = GOLD

    hdr = row(None, YEARS); 
    for col in COLS: ws[f"{col}{hdr}"].font = F_B; ws[f"{col}{hdr}"].border = Border(bottom=THIN); ws[f"{col}{hdr}"].alignment = CENTER
    R = {}
    R["rev"] = row("Revenue", [f"=RPM!{c}58" for c in RPM_COLS], FM, note="RPM row 58 (total revenue).")
    R["lab"] = row("Label %", [0.50], FP1, note="Recording-royalty rate before Discovery Mode, % of total revenue, blended across tiers. Major-label Premium headline ~52%, independents ~50%, ad tier on the greater of revenue share or per-stream minimums. Chosen so that royalties after the DM haircut land on Loud & Clear's €10.1bn (58.8% of 2025 revenue). Held flat.")
    flat(0.50, R["lab"])
    R["pub"] = row("Publisher %", [0.11], FP1, note="Mechanical + performance royalties, % of total revenue. US Premium statutory all-in 15.25% (Phonorecords IV, 2025) less the bundle discount since Mar-24; non-US nearer 10-13%. Not reduced by Discovery Mode. Held flat.")
    flat(0.11, R["pub"])
    R["roy"] = row("Royalties", [f"={c}{R['rev']}*({c}{R['lab']}+{c}{R['pub']})" for c in COLS], FM, note="Revenue x (label % + publisher %): the royalty pool before the Discovery Mode haircut.")
    blank()
    def dm_path(label, k25, k26, k27, note):
        vals = [f"='DM Inputs'!{k25}", f"='DM Inputs'!{k26}", f"='DM Inputs'!{k27}", None, None, None]
        rr = row(label, vals, FP1, note=note)
        ws[f"E{rr}"].value = f"=D{rr}+(D{rr}-C{rr})*2/3"; ws[f"F{rr}"].value = f"=E{rr}+(D{rr}-C{rr})*1/3"; ws[f"G{rr}"].value = f"=F{rr}"
        for col in ["E", "F", "G"]: ws[f"{col}{rr}"].font = F_TXT; ws[f"{col}{rr}"].alignment = CENTER; ws[f"{col}{rr}"].number_format = FP1; ws[f"{col}{rr}"].fill = GOLD
        return rr
    R["e"] = dm_path("Discovery Mode Eligible Streams", "D30", "E31", "F31", "% of streams served in Radio / Autoplay / Mixes. FY25 = DM Inputs D30; FY26-FY27 = the active case rows on DM Inputs (E31, F31); FY28 adds two thirds of the FY27 step, FY29 one third, FY30 flat.")
    R["s"] = dm_path("DM % of DM-eligible streams ", "D35", "E36", "F36", "% of eligible-context streams that are Discovery Mode streams. FY25 = DM Inputs D35; FY26-FY27 = DM Inputs E36, F36; same fade to flat by FY30.")
    R["hc"] = row("DM haircut", [0.30], FP0, note="Program term: 30% lower recording royalty on DM-context streams. Held flat; lower a later year to model a claw-back at label renewals.")
    flat(0.30, R["hc"])
    R["dmr"] = row("DM royalties", [f"={c}{R['rev']}*{c}{R['lab']}*{c}{R['e']}*{c}{R['s']}" for c in COLS], FM, note="Label royalties attributable to Discovery Mode streams before the haircut = revenue x label % x eligible share x DM share. Publishing is not discounted, so only the label line is in scope.")
    R["dmn"] = row("DM netted royalties", [f"={c}{R['dmr']}*(1-{c}{R['hc']})" for c in COLS], FM, note="What is actually paid on those streams = DM royalties x (1 - haircut).")
    blank()
    R["ab"] = row("Audiobook licensing", [f"='Audiobook GM Bridge'!D41", f"='Audiobook GM Bridge'!E41", f"='Audiobook GM Bridge'!F41", None, None, None], FM, note="FY25-FY27 from the audiobook bridge (row 41, cost paid to book publishers). FY28-FY30 grown at the rates below.")
    R["abg"] = row("   % y/y growth", ["", f"=C{R['ab']}/B{R['ab']}-1", f"=D{R['ab']}/C{R['ab']}-1", 0.20, 0.15, 0.10], FP0, italic=True, note="FY28-FY30: inputs, decelerating from the audiobook tab's FY27 rate.")
    for col, prev in zip(["E", "F", "G"], ["D", "E", "F"]):
        c = ws[f"{col}{R['ab']}"]; c.value = f"={prev}{R['ab']}*(1+{col}{R['abg']})"; c.font = F_TXT; c.alignment = CENTER; c.number_format = FM
    blank()
    R["oth"] = row("Other cost of revenue (payment fees, podcasts, delivery)", [f"=-IS!AA12-(B{R['roy']}-B{R['dmr']}+B{R['dmn']}+B{R['ab']})"] + [f"={c}{R['rev']}*$B${r[0]}/$B${R['rev']}" for c in COLS[1:]], FM,
                   note="Added so FY25 ties to reported cost of revenue (IS 2025, €11,690M): the plug is payment processing (~3% of revenue), podcast content and production (~2%), streaming delivery, hosting and support (~2%). Held at the FY25 % of revenue. Delete this row for the pure royalty waterfall; FY25 gross margin then reads ~39% instead of 32%.")
    for col in COLS[1:]: ws[f"{col}{R['oth']}"].fill = GOLD
    blank()
    R["cogs"] = row("COGS", [f"={c}{R['roy']}-{c}{R['dmr']}+{c}{R['dmn']}+{c}{R['ab']}+{c}{R['oth']}" for c in COLS], FM, bold=True, top=True, note="Royalties - DM royalties + DM netted royalties + audiobook licensing + other.")
    R["gp"] = row("Gross profit", [f"={c}{R['rev']}-{c}{R['cogs']}" for c in COLS], FM)
    R["gm"] = row("Gross margin %", [f"={c}{R['gp']}/{c}{R['rev']}" for c in COLS], FP1, bold=True, fill=YELLOW)
    blank()
    row("Memo", bold=True)
    R["rep"] = row("Reported COGS (IS)", [f"=-IS!AA12"], FM, note="FY25 20-F / IS. FY25 gross margin reproduces 32.0% by construction.")
    row("Royalties after Discovery Mode, % of revenue", [f"=({c}{R['roy']}-{c}{R['dmr']}+{c}{R['dmn']})/{c}{R['rev']}" for c in COLS], FP1, note="Check against Loud & Clear: €10.1bn paid to the music industry in 2025 = 58.8% of revenue.")
    row("Discovery Mode saving", [f"={c}{R['dmr']}-{c}{R['dmn']}" for c in COLS], FM, note="DM royalties - DM netted royalties = the cost-of-revenue reduction.")
    row("   as bp of gross margin", [f"=({c}{R['dmr']}-{c}{R['dmn']})/{c}{R['rev']}*10000" for c in COLS], FBP)
    row("Gross margin % with no Discovery Mode", [f"=({c}{R['rev']}-{c}{R['cogs']}-({c}{R['dmr']}-{c}{R['dmn']}))/{c}{R['rev']}" for c in COLS], FP1)
    row("Blue = input; gold = held flat / case-driven / fade; black = formula. After pasting into the model, Find & Replace \"[SPOT_GM_Waterfall_Tab.xlsx]\" with nothing so links point at the model's own tabs; the stub sheets in this file are not needed there.", italic=True)
    # stubs
    for name, cells in [("RPM", {"B2": "STUB", **{f"{c}2": y for c, y in zip(RPM_COLS, YEARS)}, **{f"{c}58": v for c, v in zip(RPM_COLS, RPM_TOTAL)}, "B58": "Total Revenue"}),
                        ("DM Inputs", {"B2": "STUB", **DMI}), ("Audiobook GM Bridge", {"B2": "STUB", **AB, "B41": "C: Licensing cost paid to book publishers"}),
                        ("IS", {"B2": "STUB", "AA3": "2025", "B12": "Cost of Revenue", "AA12": -11690})]:
        st = wb.create_sheet(name)
        for ref, v in cells.items(): st[ref] = v; st[ref].font = F_IN if isinstance(v, (int, float)) else F_I
    wb.save(OUT); print("saved", OUT)

if __name__ == "__main__":
    build()
