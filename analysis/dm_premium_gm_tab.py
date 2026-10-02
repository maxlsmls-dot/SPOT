#!/usr/bin/env python3
"""'DM Premium GM' tab: the Sheet2 chain, annual FY'25-FY'30, styled like the operating model.

  Premium revenue (RPM row 55)
  x normal label royalty % (before Discovery Mode)
  e = % of premium streams in DM-eligible contexts, s = % of those that are Discovery Mode   (FY'25 from DM Inputs;
      Bear/Base/Bull rows FY'26-FY'30 in the Assumptions-tab pattern, faded linearly to flat by FY'30)
  royalty % after haircut = normal royalty % x (1 - e x s x 30%)
  premium gross margin   = 1 - royalty % after haircut - other premium cost of revenue %

Output: data/SPOT_DM_PremiumGM_Tab.xlsx with the tab plus three STUB sheets (RPM, Valuation, DM Inputs) holding
only the cells the tab references, so it calculates stand-alone. After pasting the tab into the model, Find & Replace
"[SPOT_DM_PremiumGM_Tab.xlsx]" with nothing so the links point at the model's own tabs, then delete nothing else.
Run: python3 analysis/dm_premium_gm_tab.py   (then recalc with LibreOffice)
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "data", "SPOT_DM_PremiumGM_Tab.xlsx")

NAVY = PatternFill("solid", fgColor="FF002060"); BAND = PatternFill("solid", fgColor="FFE7E6E6")
GOLD = PatternFill("solid", fgColor="FFFFF2CC"); YELLOW = PatternFill("solid", fgColor="FFFFFF00")
F_HDR = Font(name="Calibri", size=11, bold=True, color="FFFFFFFF")
F_TXT = Font(name="Calibri", size=11); F_B = Font(name="Calibri", size=11, bold=True); F_I = Font(name="Calibri", size=11, italic=True)
F_IN = Font(name="Calibri", size=11, color="FF0070C0"); F_IN_B = Font(name="Calibri", size=11, bold=True, color="FF0070C0")
THIN = Side(style="thin"); CENTER = Alignment(horizontal="center"); LEFT = Alignment(horizontal="left"); NOTE = Alignment(wrap_text=True, vertical="top")
FM = '#,##0_);\\(#,##0\\)'; FBP = '0_);\\(0\\)'; FP0 = '0%'; FP1 = '0.0%'; FP2 = '0.00%'; FX = '0.00'

SH = "DM Premium GM"
YC = ["C", "D", "E", "F", "G", "H"]; YEARS = ["FY'25", "FY'26E", "FY'27E", "FY'28E", "FY'29E", "FY'30E"]
RPM_COLS = ["Q", "V", "AA", "AB", "AC", "AD"]        # FY'25 ... FY'30 on the RPM tab
SEL = "Valuation!$C$3"
def case(c, bear, base, bull): return f'=CHOOSE(MATCH({SEL}, {{"Bear","Base","Bull"}}, 0), {c}{bear}, {c}{base}, {c}{bull})'

# --- values seeded from the model (SPOT_v2.xlsx) for the stubs and the case rows
RPM_PREMIUM = [15350, 17502, 18991, 20723, 22331, 23581]     # RPM row 55, FY'25-FY'30 (model values at build time)
RPM_TOTAL = [17186, 19410, 21175, 23169, 25017, 26482]       # RPM row 58
DMI_E25, DMI_S25 = 0.512, 0.33                                # 'DM Inputs'!D30, D35
E_CASES = {"Bear": [0.515, 0.520], "Base": [0.520, 0.525], "Bull": [0.525, 0.530]}   # 'DM Inputs' E32:F34, FY'26E-FY'27E levels
S_CASES = {"Bear": [0.320, 0.320], "Base": [0.330, 0.335], "Bull": [0.340, 0.350]}   # 'DM Inputs' E37:F39

def setup(ws, title):
    ws.sheet_view.showGridLines = False
    for c, w in {"A": 8.9, "B": 62, **{c: 13 for c in YC}, "I": 78}.items(): ws.column_dimensions[c].width = w
    ws["B2"] = title
    for col in ["B"] + YC + ["I"]:
        ws[f"{col}2"].fill = NAVY; ws[f"{col}2"].font = F_HDR; ws[f"{col}3"].fill = BAND
    for col, y in zip(YC, YEARS): ws[f"{col}2"] = y; ws[f"{col}2"].alignment = CENTER
    ws["I2"] = "Notes"; ws.freeze_panes = "A4"

class W:
    def __init__(self, ws): self.ws = ws; self.r = 3
    def row(self, label=None, vals=None, note=None, bold=False, italic=False, fill=None, top=False, box=False, fmt=None, label_font=None):
        self.r += 1; ws = self.ws; rr = self.r
        if label is not None:
            c = ws.cell(row=rr, column=2, value=label); c.font = label_font or (F_B if bold else (F_I if italic else F_TXT))
            if italic: c.alignment = LEFT
        if vals:
            for col, v in zip(YC, vals):
                if v is None or v == "": continue
                c = ws[f"{col}{rr}"]; c.value = v; c.alignment = CENTER
                is_f = isinstance(v, str) and v.startswith("=")
                c.font = (F_B if bold else F_TXT) if is_f else (F_IN_B if bold else F_IN)
                if fmt: c.number_format = fmt
        if note: c = ws[f"I{rr}"]; c.value = note; c.font = F_I; c.alignment = NOTE
        if fill:
            for col in ["B"] + YC: ws[f"{col}{rr}"].fill = fill
        if top:
            for col in ["B"] + YC: ws[f"{col}{rr}"].border = Border(top=THIN)
        if box:
            ws[f"B{rr}"].border = Border(left=THIN, top=THIN, bottom=THIN)
            for col in YC: ws[f"{col}{rr}"].border = Border(top=THIN, bottom=THIN)
        return rr
    def blank(self): self.r += 1
    def section(self, label): return self.row(label, bold=True)
    def cases(self, label, fy25, cases, fmt, note):
        """Active row (CHOOSE on Valuation!$C$3) + Bear/Base/Bull rows. FY'25 = link; FY'26E-FY'27E = blue levels;
        FY'28E-FY'30E = prior year + the FY'27E increment x the fade factor (flat by FY'30E)."""
        ws = self.ws
        ra = self.row(label, [fy25, None, None, None, None, None], note, fmt=fmt)
        rb, rbase, rbu = ra + 1, ra + 2, ra + 3
        for col in YC[1:]:
            c = ws[f"{col}{ra}"]; c.value = case(col, rb, rbase, rbu); c.font = F_TXT; c.alignment = CENTER; c.number_format = fmt; c.fill = GOLD
        for name, rr, bd in [("Bear", rb, Border(left=THIN, top=THIN)), ("Base", rbase, Border(left=THIN)), ("Bull", rbu, Border(left=THIN, bottom=THIN))]:
            vals = [f"=C{ra}", cases[name][0], cases[name][1],
                    f"=E{rr}+(E{rr}-D{rr})*F${self.fade}", f"=F{rr}+(E{rr}-D{rr})*G${self.fade}", f"=G{rr}+(E{rr}-D{rr})*H${self.fade}"]
            self.row(name, vals, fmt=fmt); ws[f"B{rr}"].border = bd
            ws[f"C{rr}"].font = F_TXT
        return ra

def build():
    wb = Workbook(); wb.properties.creator = "SPOT"; wb.properties.lastModifiedBy = "SPOT"
    ws = wb.active; ws.title = SH; setup(ws, "Discovery Mode: Premium Gross Margin Build"); w = W(ws); R = {}
    w.row("€ in millions unless stated.  Premium revenue from the RPM; the normal label royalty rate is cut by 30% on the share of premium streams that are Discovery Mode streams (e x s); the royalty pool after the haircut over premium revenue, plus other premium cost of revenue, gives premium gross margin.", label_font=F_I)
    w.blank()
    w.section("A. Premium revenue (RPM)")
    R["prem"] = w.row("Premium revenue", [f"=RPM!{c}55" for c in RPM_COLS], "RPM row 55 (FY'25 actual; FY'26-FY'30 projection, case-driven through the RPM assumptions).", fmt=FM)
    R["tot"] = w.row("   Total revenue (memo)", [f"=RPM!{c}58" for c in RPM_COLS], italic=True, fmt=FM)
    R["g"] = w.row("   % y/y premium revenue growth", ["", f"=D{R['prem']}/C{R['prem']}-1", f"=E{R['prem']}/D{R['prem']}-1", f"=F{R['prem']}/E{R['prem']}-1", f"=G{R['prem']}/F{R['prem']}-1", f"=H{R['prem']}/G{R['prem']}-1"], italic=True, fmt=FP1)
    w.blank()
    w.section("B. Normal royalty rate (before Discovery Mode)")
    R["roy"] = w.row("Label and publisher royalties, % of premium revenue, before Discovery Mode", [0.61, "=C{r}", "=D{r}", "=E{r}", "=F{r}", "=G{r}"], "Input. Calibration: Loud & Clear puts music royalties at €10.1bn in 2025 (58.8% of total revenue); the premium tier carries a higher rate than the ad tier. 61% leaves ~8% of premium revenue for other cost of revenue at the reported 34% premium gross margin. Replace with the model's own royalty-rate line if one exists; held flat.", fmt=FP1)
    for col in YC[1:]: ws[f"{col}{R['roy']}"].value = f"={chr(ord(col)-1)}{R['roy']}"; ws[f"{col}{R['roy']}"].fill = GOLD
    w.blank()
    w.section("C. Discovery Mode stream shares (FY'25 from DM Inputs; cases FY'26E-FY'27E as on DM Inputs; faded to flat by FY'30E)")
    R["fade"] = w.row("Fade factor: share of the FY'27E increment added in each later year", ["", "", "", 0.667, 0.333, 0.0], "Linear fade: FY'28E adds two thirds of the FY'27E step, FY'29E one third, FY'30E nothing, so both shares are flat by FY'30E. Set all three to 0 for flat from FY'28E, or to 1 to keep the FY'27E step going.", fmt=FX)
    for col in ["F", "G", "H"]: ws[f"{col}{R['fade']}"].fill = GOLD
    w.fade = R["fade"]
    R["e"] = w.cases("e: % of premium streams served in DM-eligible contexts (Radio, Autoplay, Mixes)", "='DM Inputs'!D30", E_CASES, FP1,
                     "FY'25 links to DM Inputs D30. Bear/Base/Bull FY'26E-FY'27E are the levels on DM Inputs rows 32-34 (typed here so the tab does not depend on that tab's layout). Active row follows Valuation!C3 like the Assumptions tab.")
    R["s"] = w.cases("s: % of DM-eligible-context streams that are Discovery Mode streams", "='DM Inputs'!D35", S_CASES, FP1,
                     "FY'25 links to DM Inputs D35 (burner-account sampling). Bear/Base/Bull FY'26E-FY'27E from DM Inputs rows 37-39.")
    R["es"] = w.row("e x s: Discovery Mode streams, % of premium streams", [f"={c}{R['e']}*{c}{R['s']}" for c in YC], "The share of the royalty pool that takes the haircut.", bold=True, top=True, fmt=FP1)
    w.blank()
    w.section("D. Royalty haircut")
    R["hc"] = w.row("Discovery Mode royalty haircut", [0.30] * 6, "Program term: 30% lower royalty on streams served in DM contexts. Set a lower later-year value to model a claw-back at the major-label renewals.", fmt=FP0)
    R["royhc"] = w.row("Royalty % after Discovery Mode haircut = normal royalty % x (1 - e x s x haircut)", [f"={c}{R['roy']}*(1-{c}{R['es']}*{c}{R['hc']})" for c in YC], bold=True, top=True, fmt=FP2)
    R["pool"] = w.row("Royalty pool after haircut, €M = premium revenue x royalty % after haircut", [f"={c}{R['prem']}*{c}{R['royhc']}" for c in YC], fmt=FM)
    R["sav"] = w.row("   Discovery Mode saving, €M = premium revenue x normal royalty % x e x s x haircut", [f"={c}{R['prem']}*{c}{R['roy']}*{c}{R['es']}*{c}{R['hc']}" for c in YC], italic=True, fmt=FM)
    w.blank()
    w.section("E. Premium gross margin")
    R["gm25"] = w.row("Reported premium gross margin, FY'25 (calibration)", [0.34], "20-F FY2025: Premium gross margin 34% (33% in FY'24). Used only to size other cost of revenue in FY'25.", fmt=FP1)
    R["oth"] = w.row("Other premium cost of revenue, % of premium revenue (payment fees, audiobook and podcast content, delivery)", [f"=1-C{R['gm25']}-C{R['royhc']}", "=C{r}", "=D{r}", "=E{r}", "=F{r}", "=G{r}"],
                     "FY'25 is the plug that reproduces the reported 34%; held flat after. Audiobook cost growth is modelled on the audiobook tabs: apply that bridge's y/y headwind on top of this margin rather than raising this line.", fmt=FP2)
    for col in YC[1:]: ws[f"{col}{R['oth']}"].value = f"={chr(ord(col)-1)}{R['oth']}"; ws[f"{col}{R['oth']}"].fill = GOLD
    R["cor"] = w.row("Premium cost of revenue, €M = premium revenue x (royalty % after haircut + other %)", [f"={c}{R['prem']}*({c}{R['royhc']}+{c}{R['oth']})" for c in YC], fmt=FM)
    R["gp"] = w.row("Premium gross profit, €M", [f"={c}{R['prem']}-{c}{R['cor']}" for c in YC], fmt=FM)
    R["gm"] = w.row("Premium Gross Margin %", [f"={c}{R['gp']}/{c}{R['prem']}" for c in YC], "Equals 1 - royalty % after haircut - other %. FY'25 reproduces the reported 34% by construction. Link the IS premium gross margin (or the Model gross-margin row, weighted with ad-supported) to this row.", bold=True, fill=YELLOW, box=True, fmt=FP1)
    R["dgm"] = w.row("   y/y change, bp", ["", f"=(D{R['gm']}-C{R['gm']})*10000", f"=(E{R['gm']}-D{R['gm']})*10000", f"=(F{R['gm']}-E{R['gm']})*10000", f"=(G{R['gm']}-F{R['gm']})*10000", f"=(H{R['gm']}-G{R['gm']})*10000"], italic=True, fmt=FBP)
    w.blank()
    w.section("F. Memo: what Discovery Mode is worth inside that margin")
    R["gm0"] = w.row("Premium gross margin with no Discovery Mode = 1 - normal royalty % - other %", [f"=1-{c}{R['roy']}-{c}{R['oth']}" for c in YC], fmt=FP1)
    R["dmbp"] = w.row("Discovery Mode contribution to premium gross margin, bp = normal royalty % x e x s x haircut", [f"=({c}{R['gm']}-{c}{R['gm0']})*10000" for c in YC], bold=True, fmt=FBP)
    R["dmbpd"] = w.row("   y/y change in the contribution, bp (what the fade does)", ["", f"=D{R['dmbp']}-C{R['dmbp']}", f"=E{R['dmbp']}-D{R['dmbp']}", f"=F{R['dmbp']}-E{R['dmbp']}", f"=G{R['dmbp']}-F{R['dmbp']}", f"=H{R['dmbp']}-G{R['dmbp']}"], "Positive only while e x s is still rising; zero once the shares are flat, because the saving then grows exactly with premium revenue.", italic=True, fmt=FBP)
    R["sav2"] = w.row("Discovery Mode saving as % of total revenue, bp (consolidated-margin equivalent)", [f"={c}{R['sav']}/{c}{R['tot']}*10000" for c in YC], italic=True, fmt=FBP)
    w.blank()
    w.section("G. Memo: consolidated gross margin (for the Model / RDCF gross-margin row)")
    R["adgm"] = w.row("Ad-Supported gross margin (input; FY'25 from 20-F, held flat)", [0.18, "=C{r}", "=D{r}", "=E{r}", "=F{r}", "=G{r}"], "FY'25 20-F: Ad-Supported gross profit €330M on €1,836M = 18%. Marquee and Showcase sit here. Replace with the model's own ad-supported build.", fmt=FP1)
    for col in YC[1:]: ws[f"{col}{R['adgm']}"].value = f"={chr(ord(col)-1)}{R['adgm']}"; ws[f"{col}{R['adgm']}"].fill = GOLD
    R["adrev"] = w.row("Ad-Supported revenue (RPM row 56)", [f"=RPM!{c}56" for c in RPM_COLS], fmt=FM)
    R["cgm"] = w.row("Consolidated gross margin = (premium GP + ad-supported revenue x ad GM) / total revenue", [f"=({c}{R['gp']}+{c}{R['adrev']}*{c}{R['adgm']})/{c}{R['tot']}" for c in YC], "FY'25 reproduces ~32.3% vs 32.0% reported; the gap is rounding in the 34% and 18% segment margins.", bold=True, top=True, fmt=FP1)
    w.blank()
    w.row("Colour key: blue = hard-coded input; light gold = case-driven cell, held-flat assumption or placeholder to link to the model; yellow = output row; black = formula. Case rows under each lever follow Valuation!C3 like the Assumptions tab.", label_font=F_I)

    # ----------------------------------------------------------------- stubs so the tab calculates stand-alone
    for name, cells in [("RPM", {"B2": "STUB: mirrors RPM rows 2/55/56/58 for the FY columns. Delete after pasting the tab into the model.", **{f"{c}2": y for c, y in zip(RPM_COLS, YEARS)},
                                 **{f"{c}55": v for c, v in zip(RPM_COLS, RPM_PREMIUM)}, **{f"{c}58": v for c, v in zip(RPM_COLS, RPM_TOTAL)},
                                 **{f"{c}56": t - p for c, t, p in zip(RPM_COLS, RPM_TOTAL, RPM_PREMIUM)}, "B55": "Premium Revenue", "B56": "Ad-Supported Revenue (Free Tier)", "B58": "Total Revenue"}),
                        ("Valuation", {"B2": "Case Selector", "B3": "Toggle", "C3": "Base", "E3": "STUB: mirrors Valuation!C3. Delete after pasting."}),
                        ("DM Inputs", {"B30": "e: share of all streams served from DM-eligible surfaces", "D30": DMI_E25, "B35": "s: enrolled share of eligible-surface streams", "D35": DMI_S25,
                                       "B2": "STUB: mirrors DM Inputs D30 and D35. Delete after pasting."})]:
        st = wb.create_sheet(name)
        for ref, v in cells.items():
            st[ref] = v; st[ref].font = F_IN if isinstance(v, (int, float)) else F_I
        st.column_dimensions["B"].width = 60
    wb.save(OUT); print("saved", OUT)

if __name__ == "__main__":
    build()
