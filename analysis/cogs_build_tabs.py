#!/usr/bin/env python3
"""Three tabs for the SPOT operating model, FY'24-FY'30E:

  'Audiobook Inputs'   the audiobook assumptions, extended from FY'27E to FY'30E (revenue linked to the RPM; case selector on Valuation!C3)
  'Audiobook GM Bridge' B (bundle saving) -> C (licensing cost) -> A (paid revenue) -> N, extended to FY'30E
  'COGS Build'          the user's waterfall, corrected and completed:
                          revenue (RPM) -> label % / publisher % (with y/y glide rows) -> royalties
                          -> Discovery Mode eligible x DM share x haircut -> royalties net of the DM saving
                          -> audiobook licensing cost (bridge C) + bundle-saving hand-back vs FY25 (bridge B)
                          -> other cost of revenue (FY25 plug to reported COGS, with a y/y glide) -> COGS, GP, GM

Fixes vs the model's COGS Build: DM netted royalties now subtracts the 30% saving (was subtracting the 70% paid);
'Other' is no longer circular; audiobooks link to licensing cost (bridge row C) rather than the net contribution N.
Output: data/SPOT_COGS_Build_Tabs.xlsx + stub sheets (RPM, Valuation, DM Inputs, IS) so it calculates stand-alone.
Install: delete the model's three tabs, paste these, then Find & Replace "[SPOT_COGS_Build_Tabs.xlsx]" with nothing.
Run: python3 analysis/cogs_build_tabs.py   (then recalc with LibreOffice)
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "data", "SPOT_COGS_Build_Tabs.xlsx")
NAVY = PatternFill("solid", fgColor="FF002060"); BAND = PatternFill("solid", fgColor="FFE7E6E6")
GOLD = PatternFill("solid", fgColor="FFFFF2CC"); YELLOW = PatternFill("solid", fgColor="FFFFFF00")
F_HDR = Font(name="Calibri", size=11, bold=True, color="FFFFFFFF")
F_TXT = Font(name="Calibri", size=11); F_B = Font(name="Calibri", size=11, bold=True); F_I = Font(name="Calibri", size=11, italic=True)
F_IN = Font(name="Calibri", size=11, color="FF0070C0"); F_IN_B = Font(name="Calibri", size=11, bold=True, color="FF0070C0")
THIN = Side(style="thin"); CENTER = Alignment(horizontal="center"); LEFT = Alignment(horizontal="left"); NOTE = Alignment(wrap_text=True, vertical="top")
FM = '#,##0_);\\(#,##0\\)'; FBP = '0_);\\(0\\)'; FP0 = '0%'; FP1 = '0.0%'; FP2 = '0.00%'; F0 = '0'; FX = '0.00'; FRATE = '0.00000'; FDELTA = '+0;-0;0'
SEL = "Valuation!$C$3"
SH_IN, SH_BR, SH_CB = "Audiobook Inputs", "Audiobook GM Bridge", "COGS Build"
IN = f"'{SH_IN}'!"; BRQ = f"'{SH_BR}'!"
YC = ["C", "D", "E", "F", "G", "H", "I"]; YEARS = ["FY'24", "FY'25", "FY'26E", "FY'27E", "FY'28E", "FY'29E", "FY'30E"]
RPM_COLS = ["L", "Q", "V", "AA", "AB", "AC", "AD"]
RPM_TOTAL = [15673, 17186, 19410, 21172, 23166, 25017, 26482]      # model values at build time (stub only)
DMI = {"D30": 0.512, "E31": 0.52, "F31": 0.525, "G31": 0.5275, "H31": 0.529, "I31": 0.53,
       "D35": 0.33, "E36": 0.33, "F36": 0.335, "G36": 0.337, "H36": 0.339, "I36": 0.34}
SP = "J"; NC = "K"   # spare / notes columns on the audiobook tabs

def setup(ws, title, years=YEARS, ycols=YC, notes=NC, spare=(SP,), label_w=58):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 8.9; ws.column_dimensions["B"].width = label_w
    for c in list(ycols) + list(spare): ws.column_dimensions[c].width = 13
    ws.column_dimensions[notes].width = 78
    ws["B2"] = title
    for col in ["B"] + list(ycols) + list(spare) + [notes]:
        ws[f"{col}2"].fill = NAVY; ws[f"{col}2"].font = F_HDR; ws[f"{col}3"].fill = BAND
    for col, y in zip(ycols, years): ws[f"{col}2"] = y; ws[f"{col}2"].alignment = CENTER
    ws[f"{notes}2"] = "Notes"; ws.freeze_panes = "A4"

class Writer:
    def __init__(self, ws, ycols=YC, notes=NC, spare=SP): self.ws = ws; self.r = 3; self.yc = ycols; self.nc = notes; self.sp = spare
    def row(self, label=None, vals=None, note=None, bold=False, italic=False, fill=None, top=False, box=False, g=None, fmt=None, label_font=None, fmts=None):
        self.r += 1; ws = self.ws; rr = self.r
        if label is not None:
            c = ws.cell(row=rr, column=2, value=label); c.font = label_font or (F_B if bold else (F_I if italic else F_TXT))
            if italic: c.alignment = LEFT
        if vals:
            for i, (col, v) in enumerate(zip(self.yc, vals)):
                if v is None or v == "": continue
                self._cell(f"{col}{rr}", v, bold, (fmts[i] if fmts else fmt))
        if g is not None: self._cell(f"{self.sp}{rr}", g, bold, fmt)
        if note: c = ws[f"{self.nc}{rr}"]; c.value = note; c.font = F_I; c.alignment = NOTE
        if fill:
            for col in ["B"] + list(self.yc): ws[f"{col}{rr}"].fill = fill
        if top:
            for col in ["B"] + list(self.yc): ws[f"{col}{rr}"].border = Border(top=THIN)
        if box:
            ws[f"B{rr}"].border = Border(left=THIN, top=THIN, bottom=THIN)
            for col in self.yc: ws[f"{col}{rr}"].border = Border(top=THIN, bottom=THIN)
        return rr
    def _cell(self, ref, v, bold, fmt):
        c = self.ws[ref]; c.value = v; c.alignment = CENTER
        is_f = isinstance(v, str) and v.startswith("=")
        c.font = (F_B if bold else F_TXT) if is_f else (F_IN_B if bold else F_IN)
        if fmt: c.number_format = fmt
    def blank(self): self.r += 1
    def section(self, label): return self.row(label, bold=True)
    def subhdr(self, labels, g=None):
        rr = self.row(labels[0], label_font=F_I)
        for col, t in zip(self.yc, labels[1:]):
            c = self.ws[f"{col}{rr}"]; c.value = t; c.font = F_B; c.alignment = CENTER; c.border = Border(bottom=THIN)
        if g: c = self.ws[f"{self.sp}{rr}"]; c.value = g; c.font = F_B; c.alignment = CENTER; c.border = Border(bottom=THIN)
        return rr

def build():
    wb = Workbook(); wb.properties.creator = "SPOT"; wb.properties.lastModifiedBy = "SPOT"
    A = {}
    # ======================================================================= Audiobook Inputs
    ws = wb.active; ws.title = SH_IN; setup(ws, "Audiobook Assumptions"); w = Writer(ws)
    rr = w.row("Case selector", [f"={SEL}"], "Linked to Valuation!C3. Drives the licensing-cost growth assumption in section E. Bear = adverse for SPOT (faster cost growth).", fmt=None)
    ws[f"C{rr}"].fill = GOLD; A["case"] = f"{IN}$C${rr}"; sel = A["case"]
    w.blank()
    w.section("A. Group revenue, €M (denominator for basis points)")
    rr = w.row("Total revenue", [f"=RPM!{c}58" for c in RPM_COLS], "Linked to RPM row 58 (FY'24-FY'25 actual, FY'26-FY'30 projection).", fmt=FM)
    for col in YC: ws[f"{col}{rr}"].fill = GOLD
    for col, k in zip(YC, ["rev24", "rev25", "rev26", "rev27", "rev28", "rev29", "rev30"]): A[k] = f"{IN}${col}${rr}"
    w.blank()
    w.section("B. Bundle royalty saving: MLC contingency disclosed in Spotify filings, €M cumulative from 1 Mar 2024")
    series = [("mlc_cum_q224", 46, "6-K Q2-24: cumulative to 30 Jun 2024", "6-K Q2-24 legal proceedings note: 'approximately €46 million, of which approximately €35 million relates to the three months ended June 30, 2024'."),
              ("mlc_q224_only", 35, "   of which Apr-Jun 2024", "Same sentence; implies March 2024 alone was about €11M."),
              ("mlc_cum_q324", 94, "6-K Q3-24: cumulative to 30 Sep 2024", "6-K Q3-24 legal proceedings note."),
              ("mlc_cum_q424", 150, "20-F FY24: cumulative to 31 Dec 2024", "Form 20-F FY2024, legal proceedings."),
              ("mlc_cum_q125", 205, "6-K Q1-25: cumulative to 31 Mar 2025", "6-K Q1-25 legal proceedings note."),
              ("mlc_cum_q225", 256, "6-K Q2-25: cumulative to 30 Jun 2025", "6-K Q2-25 legal proceedings note."),
              ("mlc_cum_q325", 308, "6-K Q3-25: cumulative to 30 Sep 2025", "6-K Q3-25: first filing to add 'any liability would be partially offset by direct deals with publishers'."),
              ("mlc_cum_q425", 358, "20-F FY25: cumulative to 31 Dec 2025", "Form 20-F FY2025, legal proceedings."),
              ("mlc_cum_q126", 410, "6-K Q1-26: cumulative to 31 Mar 2026", "6-K Q1-26 (about $473M at the time)."),
              ("mlc_cum_q226", 437, "6-K Q2-26: cumulative to 30 Jun 2026 (verify)", "Quoted by Music Ally, 7 Sep 2026, as €437M; a second reading gives €473M. Verify in the Q2-26 6-K before quoting.")]
    for k, v, label, note in series:
        rr = w.row(label, [v], note, fmt=FM); A[k] = f"{IN}$C${rr}"
    rr = w.row("Gross bundle saving, forecast placeholders", ["", "", 205, 205, 205, 205, 205],
               "FY'26E: 4 x the average quarterly increment Q2-25 to Q1-26 (€51M). FY'27E-FY'30E: held flat; the Phonorecords IV settlement that permits the bundle rate runs to end-2027 and Phonorecords V is open, so flat is the neutral assumption.", fmt=FM)
    for col in YC[2:]: ws[f"{col}{rr}"].fill = GOLD
    for col, k in zip(YC[2:], ["brr26", "brr27", "brr28", "brr29", "brr30"]): A[k] = f"{IN}${col}${rr}"
    w.blank()
    w.section("C. Bundle saving cross-check: Billboard analysis of Spotify's MLC reports, paid tiers")
    for k, v, label, note, fmt in [("bb_q423_mech", 97.3, "Mechanical royalties paid to the MLC, Q4-23, $M", "Billboard, 'How Much Have Spotify Bundles Decreased the Mechanical Per-Stream Rate? (Analysis)'.", '0.0'),
                                   ("bb_q425_mech", 53.3, "Mechanical royalties paid to the MLC, Q4-25, $M", "Same analysis; down 45%.", '0.0'),
                                   ("bb_q423_rate", 0.00068, "Blended per-stream mechanical rate, Q4-23, $", "Same analysis.", FRATE),
                                   ("bb_q425_rate", 0.00033, "Blended per-stream mechanical rate, Q4-25, $", "Same analysis; down 51%.", FRATE),
                                   ("eurusd_q425", 1.16, "EUR/USD, Q4-25 average", "Approximate.", FX)]:
        rr = w.row(label, [v], note, fmt=fmt); A[k] = f"{IN}$C${rr}"
    w.blank()
    w.section("D. Direct publisher deals that hand part of the saving back (share of US mechanicals in FY'24 column; start month in FY'25 column)")
    pubs = [("share_umpg", "start_umpg", 0.22, 2, "Universal Music Publishing Group", "Announced 26 Jan 2025 (MBW). Share: estimate (20-25%)."),
            ("share_wcm", "start_wcm", 0.13, 3, "Warner Chappell Music", "Announced 6 Feb 2025 (MBW). Share: estimate (11-15%)."),
            ("share_kobalt", "start_kobalt", 0.05, 9, "Kobalt", "Announced 13 Aug 2025 (CMU, MBW). Share: estimate (4-7%)."),
            ("share_smp", "start_smp", 0.26, 10, "Sony Music Publishing", "Announced 18 Sep 2025 (MBW, Variety). Share: estimate (23-28%)."),
            ("share_bmg", "start_bmg", 0.03, 11, "BMG", "Announced 9 Oct 2025 (CMU). Share: estimate (2-4%).")]
    PUB = []
    for sk, st, sv, mv, label, note in pubs:
        rr = w.row(label, [sv, mv], note); ws[f"C{rr}"].number_format = FP0; ws[f"D{rr}"].number_format = F0
        A[sk] = f"{IN}$C${rr}"; A[st] = f"{IN}$D${rr}"; PUB.append((label, sk, st))
    rr = w.row("Restoration of the bundle discount in the direct deals", [0.75], "Estimate; range 50-100%.", fmt=FP0); A["restoration"] = f"{IN}$C${rr}"
    rr = w.row("MLC amended complaint succeeds in FY'27 (1 = yes)", [0], "Scenario switch: from FY'27E the independents' share of the saving goes to zero.", fmt=F0); A["mlc_loss"] = f"{IN}$C${rr}"
    w.blank()
    w.section("E. Audiobook licensing cost paid to book publishers (included and paid hours)")
    rr = w.row("FY'24 cost anchor, €M", [190], "Spotify via Axios, 15 Oct 2024: 'hundreds of millions of dollars a year'; read as ~$205M at 1.08.", fmt=FM); A["c_2024"] = f"{IN}$C${rr}"
    rr_g = w.row("Licensing cost growth, % y/y", ["", 0.58, None, None, None, None, None],
                 "FY'25: hours +37% like-for-like (Spotify newsroom Oct 2025) x ~1.15 for the wider paying base. FY'26E: Investor Day May 2026, listeners +60%, cost at ~0.7x. FY'27E-FY'30E: deceleration toward revenue growth as the hour cap binds and the listener base matures. Scenario rows below; Bear = faster cost growth.", fmt=FP1)
    bear, base, bull = rr_g + 1, rr_g + 2, rr_g + 3
    for col in YC[2:]:
        c = ws[f"{col}{rr_g}"]; c.value = f'=IF({sel}="Bear",{col}{bear},IF({sel}="Bull",{col}{bull},{col}{base}))'; c.font = F_TXT; c.alignment = CENTER; c.number_format = FP1; c.fill = GOLD
    for name, vals, bd in [("Bear", [0.60, 0.45, 0.30, 0.22, 0.15], Border(left=THIN, top=THIN)), ("Base", [0.43, 0.30, 0.20, 0.14, 0.10], Border(left=THIN)), ("Bull", [0.30, 0.20, 0.15, 0.10, 0.07], Border(left=THIN, bottom=THIN))]:
        r2 = w.row(name, ["", ""] + vals, fmt=FP1); ws[f"B{r2}"].border = bd
    for col, k in zip(YC[1:], ["cg25", "cg26", "cg27", "cg28", "cg29", "cg30"]): A[k] = f"{IN}${col}${rr_g}"
    w.blank()
    w.section("F. Paid audiobook revenue, €M (Audiobooks+, top-ups, a la carte)")
    rr = w.row("Paid audiobook revenue", [5, 30, 75, 120, 165, 210, 250], "FY'24-FY'25 estimates; FY'26E: $100M ARR reached Jul-26; FY'27E +60%; FY'28E-FY'30E +38%, +27%, +19% as payer growth decelerates. Placeholders; link to model.", fmt=FM)
    for col in YC[2:]: ws[f"{col}{rr}"].fill = GOLD
    for col, k in zip(YC, ["a24", "a25", "a26", "a27", "a28", "a29", "a30"]): A[k] = f"{IN}${col}${rr}"
    w.row("   Context: Audiobooks+ ARR, $M (Jul-26) and paying users, M (May-26)", [100, 1.0], "Investor Day 21 May 2026 via Publishers Weekly; Q2-26 call: $100M ARR reached.", fmt='0.0')
    w.blank()
    w.section("G. Consensus gross-margin expansion, bp (context only)")
    rr = w.row("Consensus y/y gross-margin expansion", ["", "", 115, 130, 100, 80, 60], "Placeholders. Link to the consensus tab.", fmt=FBP)
    for col in YC[2:]: ws[f"{col}{rr}"].fill = GOLD
    for col, k in zip(YC[2:], ["cons26", "cons27", "cons28", "cons29", "cons30"]): A[k] = f"{IN}${col}${rr}"
    w.blank()
    w.section("H. Memo, excluded from the bridge")
    rr = w.row("June 2024 US price increase, annualised revenue, €M", [500], "Estimate; a pricing effect, stays outside N.", fmt=FM); A["p_2024"] = f"{IN}$C${rr}"
    w.blank()
    w.row("Colour key: blue = hard-coded input (notes give source or derivation); light gold = forecast assumption, case-driven cell or link to the model; black = formula.", label_font=F_I)

    # ======================================================================= Audiobook GM Bridge
    bs = wb.create_sheet(SH_BR); setup(bs, "Audiobook Gross Margin Bridge"); b = Writer(bs); R = {}
    b.row("€ in millions unless stated.  N = B + A - C:  B = bundle royalty saving retained,  A = paid audiobook revenue,  C = audiobook licensing cost paid to book publishers.", label_font=F_I)
    b.blank()
    b.section("1. Bundle Royalty Saving (B)")
    b.subhdr(["MLC contingency by filing, €M", "Cumulative", "Increment", "Months", "€M / month"])
    keys = [("6-K Q2-24 (30 Jun 2024)", "mlc_cum_q224", 4), ("6-K Q3-24 (30 Sep 2024)", "mlc_cum_q324", 3), ("20-F FY24 (31 Dec 2024)", "mlc_cum_q424", 3),
            ("6-K Q1-25 (31 Mar 2025)", "mlc_cum_q125", 3), ("6-K Q2-25 (30 Jun 2025)", "mlc_cum_q225", 3), ("6-K Q3-25 (30 Sep 2025)", "mlc_cum_q325", 3),
            ("20-F FY25 (31 Dec 2025)", "mlc_cum_q425", 3), ("6-K Q1-26 (31 Mar 2026)", "mlc_cum_q126", 3), ("6-K Q2-26 (30 Jun 2026): verify", "mlc_cum_q226", 3)]
    M = {}; prev = None
    for label, k, months in keys:
        rr = b.r + 1; inc = f"=C{rr}" if prev is None else f"=C{rr}-C{prev}"
        b.row(label, [f"={A[k]}", inc, months, f"=D{rr}/E{rr}"], fmt=FM); bs[f"E{rr}"].number_format = F0; bs[f"F{rr}"].number_format = '0.0'; bs[f"E{rr}"].font = F_IN
        M[k] = rr; prev = rr
    rr = b.r + 1; b.row("   of which Apr-Jun 2024 (March 2024 alone about €11M)", [f"={A['mlc_q224_only']}", "", 3, f"=C{rr}/E{rr}"], fmt=FM); bs[f"E{rr}"].number_format = F0; bs[f"F{rr}"].number_format = '0.0'; bs[f"E{rr}"].font = F_IN
    rr = b.row("Average quarterly increment, Q2-25 to Q1-26 (x4 annualised in next column)", [f"=AVERAGE(D{M['mlc_cum_q225']}:D{M['mlc_cum_q126']})", f"=C{b.r+1}*4"], "Increments ran €50-56M a quarter with no step down after the direct deals, so the figure is gross.", fmt='0.0'); bs[f"D{rr}"].number_format = FM
    b.blank()
    b.subhdr(["Billboard cross-check (paid tiers)", "Q4-23", "Q4-25", "", "", "", "", ""])
    rb1 = b.row("Mechanical royalties paid to the MLC, $M", [f"={A['bb_q423_mech']}", f"={A['bb_q425_mech']}"], fmt='0.0')
    rb2 = b.row("Blended per-stream rate, $", [f"={A['bb_q423_rate']}", f"={A['bb_q425_rate']}"], fmt=FRATE)
    rb3 = b.row("Implied streams, bn", [f"=C{rb1}/C{rb2}/1000", f"=D{rb1}/D{rb2}/1000"], fmt='0.0')
    b.row("Implied Q4-25 discount: streams x Q4-23 rate less paid, $M (col D); in €M (col E) vs 20-F FY25 increment (col F)",
          ["", f"=D{rb3}*C{rb2}*1000-D{rb1}", f"=D{b.r+1}/{A['eurusd_q425']}", f"=D{M['mlc_cum_q425']}"], "Two independent sources agree on a gross discount of about €50M a quarter.", fmt='0.0')
    b.blank()
    b.subhdr(["Gross saving, hand-back, B"] + YEARS)
    R["gross"] = b.row("Gross bundle saving", [f"={A['mlc_cum_q424']}", f"={A['mlc_cum_q425']}-{A['mlc_cum_q424']}", f"={A['brr26']}", f"={A['brr27']}", f"={A['brr28']}", f"={A['brr29']}", f"={A['brr30']}"], "FY'24 and FY'25 from the filings; FY'26E-FY'30E placeholders on the inputs tab.", bold=True, top=True, fmt=FM)
    b.subhdr(["Months on direct licence by publisher", "", "", "", "", "", "", ""], g="Share")
    prow = {}
    for label, sk, st in PUB:
        rr = b.row("   " + label, [0, f"=13-{A[st]}", 12, 12, 12, 12, 12], g=f"={A[sk]}", fmt=F0); bs[f"{SP}{rr}"].number_format = FP0; bs[f"C{rr}"].font = F_IN; prow[sk] = rr
    pf, pl = min(prow.values()), max(prow.values())
    R["S"] = b.row("   Sum of shares on direct licences (independents are the remainder)", g=f"=SUM({SP}{pf}:{SP}{pl})", fmt=FP0)
    R["rest"] = b.row("   Restoration of the discount in the direct deals", g=f"={A['restoration']}", fmt=FP0)
    R["hb"] = b.row("Hand-back = gross x sum(share x months / 12) x restoration", [f"={c}{R['gross']}*SUMPRODUCT(${SP}${pf}:${SP}${pl},{c}{pf}:{c}{pl})/12*${SP}${R['rest']}" for c in YC], fmt=FM)
    R["B"] = b.row("B: Net bundle saving retained by Spotify", [f"={c}{R['gross']}-{c}{R['hb']}" for c in YC[:3]] + [f"=IF({A['mlc_loss']}=1,{c}{R['gross']}*${SP}${R['S']}*(1-${SP}${R['rest']}),{c}{R['gross']}-{c}{R['hb']})" for c in YC[3:]],
                   "From FY'27E: if the MLC switch on the inputs tab is 1, the independents' share of the saving goes to zero.", bold=True, top=True, fmt=FM)
    b.blank()
    b.section("2. Audiobook Licensing Cost (C)")
    R["g"] = b.row("% y/y Growth", [""] + [f"={A[k]}" for k in ["cg25", "cg26", "cg27", "cg28", "cg29", "cg30"]], "Follows the case selector on the inputs tab.", italic=True, fmt=FP1)
    rc = b.r + 1
    R["C"] = b.row("C: Licensing cost paid to book publishers", [f"={A['c_2024']}"] + [f"={p}{rc}*(1+{c}{R['g']})" for p, c in zip(YC[:-1], YC[1:])], "FY'24 anchored on Spotify's 'hundreds of millions of dollars a year'; grown by consumption.", bold=True, top=True, fmt=FM)
    b.blank()
    b.section("3. Paid Audiobook Revenue (A)")
    R["A"] = b.row("A: Paid audiobook revenue", [f"={A[k]}" for k in ["a24", "a25", "a26", "a27", "a28", "a29", "a30"]], "Counted in full: C is grown on total consumption, so it already carries the publisher cost of paid hours.", bold=True, top=True, fmt=FM)
    b.blank()
    b.section("4. Net Contribution to Gross Profit and Gross Margin")
    R["B2"] = b.row("B: Net bundle saving retained", [f"={c}{R['B']}" for c in YC], fmt=FM)
    R["A2"] = b.row("A: Paid audiobook revenue", [f"={c}{R['A']}" for c in YC], fmt=FM)
    R["C2"] = b.row("Less C: Licensing cost paid to book publishers", [f"=-{c}{R['C']}" for c in YC], fmt=FM)
    R["N"] = b.row("Audiobook Contribution to Gross Profit (N)", [f"={c}{R['B2']}+{c}{R['A2']}+{c}{R['C2']}" for c in YC], "Net effect; the COGS Build uses the components (C and B) directly rather than this line.", bold=True, fill=YELLOW, box=True, fmt=FM)
    R["rev"] = b.row("Total Revenue", [f"={A[k]}" for k in ["rev24", "rev25", "rev26", "rev27", "rev28", "rev29", "rev30"]], fmt=FM)
    R["Nbp"] = b.row("% Contribution to Gross Margin, bp", [f"={c}{R['N']}/{c}{R['rev']}*10000" for c in YC], italic=True, fmt=FBP)
    R["dN"] = b.row("Y/y Change in Contribution, €M", [""] + [f"={c}{R['N']}-{p}{R['N']}" for p, c in zip(YC[:-1], YC[1:])], fmt=FM)
    R["dNbp"] = b.row("Y/y Gross Margin Headwind, bp", [""] + [f"={c}{R['dN']}/{c}{R['rev']}*10000" for c in YC[1:]], "Negative = drag.", bold=True, fill=YELLOW, box=True, fmt=FBP)
    b.row("   of which music publishers reclaim the bundle saving (change in B)", [""] + [f"={c}{R['B2']}-{p}{R['B2']}" for p, c in zip(YC[:-1], YC[1:])], italic=True, fmt=FM)
    b.row("   of which book publishers paid on consumption (change in C)", [""] + [f"={c}{R['C2']}-{p}{R['C2']}" for p, c in zip(YC[:-1], YC[1:])], italic=True, fmt=FM)
    b.row("   of which paid audiobook revenue (change in A)", [""] + [f"={c}{R['A2']}-{p}{R['A2']}" for p, c in zip(YC[:-1], YC[1:])], italic=True, fmt=FM)
    R["cons"] = b.row("Consensus y/y gross-margin expansion, bp (context)", ["", ""] + [f"={A[k]}" for k in ["cons26", "cons27", "cons28", "cons29", "cons30"]], fmt=FBP)
    b.row("Headwind as % of consensus expansion (context)", ["", ""] + [f"=IF({c}{R['cons']}=0,0,-{c}{R['dNbp']}/{c}{R['cons']})" for c in YC[2:]], italic=True, fmt=FP0)
    b.row("Memo, excluded: June 2024 US price increase, annualised €M", [f"={A['p_2024']}"], fmt=FM)
    b.blank()
    b.section("5. Sensitivity: growth assigned to audiobook licensing cost (applied to FY'26E and FY'27E; B, A and revenue held)")
    b.subhdr(["Licensing cost growth p.a.", "N FY'26E, €M", "N FY'27E, €M", "Headwind FY'26E, bp", "Headwind FY'27E, bp", "", "", ""])
    C25 = f"$D${R['C']}"; N25 = f"$D${R['N']}"; B26, B27 = f"$E${R['B']}", f"$F${R['B']}"; A26, A27 = f"$E${R['A']}", f"$F${R['A']}"; V26, V27 = f"$E${R['rev']}", f"$F${R['rev']}"
    for g in [0.20, 0.30, 0.43, 0.50, 0.60]:
        rr = b.r + 1
        b.row(None, [g, f"={B26}+{A26}-{C25}*(1+C{rr})", f"={B27}+{A27}-{C25}*(1+C{rr})^2", f"=(D{rr}-{N25})/{V26}*10000", f"=(E{rr}-D{rr})/{V27}*10000"], fmt=FM)
        bs[f"C{rr}"].number_format = FP0; bs[f"F{rr}"].number_format = FBP; bs[f"G{rr}"].number_format = FBP
        bs.cell(row=rr, column=2, value="   Base growth, FY'26E" if abs(g - 0.43) < 1e-9 else None).font = F_I
    b.row("   Bridge base case", ["", f"=E{R['N']}", f"=F{R['N']}", f"=E{R['dNbp']}", f"=F{R['dNbp']}"], bold=True, top=True, fmt=FM)
    bs[f"F{b.r}"].number_format = FBP; bs[f"G{b.r}"].number_format = FBP

    # ======================================================================= COGS Build (user's layout: labels in A, FY25-FY30 in B-G)
    cb = wb.create_sheet(SH_CB); cb.sheet_view.showGridLines = False
    cb.column_dimensions["A"].width = 46
    CC = ["B", "C", "D", "E", "F", "G"]; CY = ["FY25", "FY26", "FY27", "FY28", "FY29", "FY30"]
    for c in CC: cb.column_dimensions[c].width = 11
    cb.column_dimensions["H"].width = 84
    rr_ = [1]; K = {}
    def row(label, vals=None, fmt=None, bold=False, italic=False, note=None, top=False, fill=None, fmts=None):
        r = rr_[0]
        if label is not None:
            c = cb[f"A{r}"]; c.value = label; c.font = F_B if bold else (F_I if italic else F_TXT)
        if vals:
            for i, (col, v) in enumerate(zip(CC, vals)):
                if v is None or v == "": continue
                c = cb[f"{col}{r}"]; c.value = v; c.alignment = CENTER
                c.font = (F_B if bold else F_TXT) if (isinstance(v, str) and v.startswith("=")) else (F_IN_B if bold else F_IN)
                f = fmts[i] if fmts else fmt
                if f: c.number_format = f
        if note: n = cb[f"H{r}"]; n.value = note; n.font = F_I; n.alignment = NOTE
        if top:
            for col in ["A"] + CC: cb[f"{col}{r}"].border = Border(top=THIN)
        if fill:
            for col in CC: cb[f"{col}{r}"].fill = fill
        rr_[0] += 1; return r
    def blank(): rr_[0] += 1
    def glide(label, start, deltas, note, lvl_fmt=FP1):
        """Level row (FY25 blue input; later years = prior + y/y bp change) and a y/y bp input row beneath."""
        lvl = row(label, [start, None, None, None, None, None], lvl_fmt, note=note)
        dl = row("   y/y change, bp", [""] + deltas, FDELTA, italic=True)
        for col in CC[1:]: cb[f"{col}{dl}"].fill = GOLD
        for p, c in zip(CC[:-1], CC[1:]):
            cell = cb[f"{c}{lvl}"]; cell.value = f"={p}{lvl}+{c}{dl}/10000"; cell.font = F_TXT; cell.alignment = CENTER; cell.number_format = lvl_fmt
        return lvl, dl
    AB = {"C": R["C"], "B": R["B"]}; BC = ["D", "E", "F", "G", "H", "I"]   # bridge columns FY'25..FY'30E
    h = row(None, CY)
    for col in CC: cb[f"{col}{h}"].font = F_B; cb[f"{col}{h}"].border = Border(bottom=THIN); cb[f"{col}{h}"].alignment = CENTER
    cb["H1"] = "Notes"; cb["H1"].font = F_B
    K["rev"] = row("Revenue", [f"=RPM!{c}58" for c in RPM_COLS[1:]], FM, note="RPM row 58.")
    K["lab"], K["labd"] = glide("Label %", 0.48, [-120, -70, -35, -20, -10], "Recording-royalty rate before the Discovery Mode haircut, % of total revenue. FY25 input. Later years glide down by the bp in the row below: price increases lift revenue faster than the per-subscriber minimums in the label deals, so the effective rate converges on the headline %; the steps shrink as the Sep-25 and Feb-26 hikes lap. This is the main source of margin expansion and the row to tune.")
    K["pub"], K["pubd"] = glide("Publisher %", 0.11, [5, 5, 0, 0, 0], "Mechanical + performance royalties, % of total revenue, after the bundle discount. Phonorecords IV escalator adds ~5 bp a year through FY27; flat after. The bundle hand-back is a separate line below.")
    K["roy"] = row("Royalties", [f"=({c}{K['lab']}+{c}{K['pub']})*{c}{K['rev']}" for c in CC], FM, note="(Label % + publisher %) x revenue, before the Discovery Mode haircut.")
    blank()
    K["e"] = row("Discovery Mode Eligible Streams", [f"='DM Inputs'!D30", f"='DM Inputs'!E31", f"='DM Inputs'!F31", f"='DM Inputs'!G31", f"='DM Inputs'!H31", f"='DM Inputs'!I31"], FP1, note="DM Inputs: FY25 D30, FY26-FY30 the active case rows (row 31).")
    K["s"] = row("DM % of DM-eligible streams", [f"='DM Inputs'!D35", f"='DM Inputs'!E36", f"='DM Inputs'!F36", f"='DM Inputs'!G36", f"='DM Inputs'!H36", f"='DM Inputs'!I36"], FP1, note="DM Inputs: FY25 D35, FY26-FY30 row 36.")
    K["hc"] = row("DM haircut", [0.30, "=B{r}", "=C{r}", "=D{r}", "=E{r}", "=F{r}"], FP0, note="30% lower recording royalty on DM-context streams; held flat.")
    for p, c in zip(CC[:-1], CC[1:]): cb[f"{c}{K['hc']}"].value = f"={p}{K['hc']}"; cb[f"{c}{K['hc']}"].font = F_TXT; cb[f"{c}{K['hc']}"].fill = GOLD
    K["dmr"] = row("DM royalties", [f"={c}{K['rev']}*{c}{K['lab']}*{c}{K['e']}*{c}{K['s']}" for c in CC], FM, note="Label royalties on Discovery Mode streams before the haircut = revenue x label % x eligible share x DM share. Publishing is not discounted.")
    K["dms"] = row("   DM saving", [f"={c}{K['dmr']}*{c}{K['hc']}" for c in CC], FM, italic=True, note="DM royalties x haircut: the cost-of-revenue reduction.")
    K["net"] = row("DM netted royalties", [f"={c}{K['roy']}-{c}{K['dms']}" for c in CC], FM, bold=True, top=True, note="Royalties less the DM saving = royalties actually paid. (Was royalties less 70% of DM royalties, which removed the paid part instead of the saving and understated royalties by ~€560M in FY25.)")
    blank()
    K["ab"] = row("Audiobook licensing", [f"={BRQ}{c}{AB['C']}" for c in BC], FM, note="Audiobook GM Bridge, C: licensing cost paid to book publishers, now extended to FY'30E. (Was the bridge's net contribution N, which nets paid audiobook revenue and the bundle saving into a cost line.)")
    K["hb"] = row("Bundle saving hand-back vs FY25", [f"={BRQ}$D${AB['B']}-{BRQ}{c}{AB['B']}" for c in BC], FM, note="Audiobook GM Bridge, B: the FY25 net bundle saving is already inside the 11% publisher rate; as the direct deals (and any MLC outcome) hand it back, publisher cost rises by FY25 B less the year's B.")
    blank()
    K["oth"], K["othd"] = glide("Other, as a % of revenue", f"=(-IS!AA12-B{K['net']}-B{K['ab']}-B{K['hb']})/B{K['rev']}", [-40, -25, -15, -10, -5],
                                 "FY25 = plug to reported cost of revenue (IS 2025, €11,690M): payment processing (~3%), podcast content and production (~2%), streaming delivery, hosting and support, Partner Program. Glides down with scale by the bp below; payment fees stay proportional, the rest does not.", lvl_fmt=FP2)
    K["othm"] = row("Other", [f"={c}{K['oth']}*{c}{K['rev']}" for c in CC], FM, note="Other % x revenue. (Was circular: % = €/revenue while € = % x revenue.)")
    blank()
    K["cogs"] = row("COGS", [f"={c}{K['net']}+{c}{K['ab']}+{c}{K['hb']}+{c}{K['othm']}" for c in CC], FM, bold=True, top=True, note="DM netted royalties + audiobook licensing + bundle hand-back + other.")
    K["gp"] = row("GP", [f"={c}{K['rev']}-{c}{K['cogs']}" for c in CC], FM)
    K["gm"] = row("GM %", [f"={c}{K['gp']}/{c}{K['rev']}" for c in CC], FP1, bold=True, fill=YELLOW, note="FY25 reproduces the reported 32.0% by construction.")
    K["dgm"] = row("   y/y change, bp", [""] + [f"=({c}{K['gm']}-{p}{K['gm']})*10000" for p, c in zip(CC[:-1], CC[1:])], FBP, italic=True)
    blank()
    row("Memo: what moves gross margin each year, bp", bold=True)
    row("   Label rate glide", [""] + [f"=-{c}{K['labd']}" for c in CC[1:]], FBP, italic=True)
    row("   Publisher rate", [""] + [f"=-{c}{K['pubd']}" for c in CC[1:]], FBP, italic=True)
    row("   Discovery Mode saving", [""] + [f"=({c}{K['dms']}/{c}{K['rev']}-{p}{K['dms']}/{p}{K['rev']})*10000" for p, c in zip(CC[:-1], CC[1:])], FBP, italic=True)
    row("   Audiobook licensing", [""] + [f"=-({c}{K['ab']}/{c}{K['rev']}-{p}{K['ab']}/{p}{K['rev']})*10000" for p, c in zip(CC[:-1], CC[1:])], FBP, italic=True)
    row("   Bundle hand-back", [""] + [f"=-({c}{K['hb']}/{c}{K['rev']}-{p}{K['hb']}/{p}{K['rev']})*10000" for p, c in zip(CC[:-1], CC[1:])], FBP, italic=True)
    row("   Other cost of revenue glide", [""] + [f"=-{c}{K['othd']}" for c in CC[1:]], FBP, italic=True)
    row("   Check: sum = y/y change", [""] + [f"=SUM({c}{rr_[0]-6}:{c}{rr_[0]-1})-{c}{K['dgm']}" for c in CC[1:]], FBP, italic=True, note="Should be ~0 (rounding from the interaction of label % and DM share).")
    row("Reported COGS FY25 (IS)", [f"=-IS!AA12"], FM, note="IS 2025 cost of revenue.")
    row("Royalties after Discovery Mode, % of revenue", [f"={c}{K['net']}/{c}{K['rev']}" for c in CC], FP1, note="Loud & Clear: €10.1bn paid to the music industry in 2025 = 58.8% of revenue; at 48% + 11% before the haircut this reads 56.6%, so ~2 points of music cost sit in 'Other'. Raise label % toward 50% to close it.")
    row("Blue = input; gold = held flat / glide; black = formula. After pasting, Find & Replace \"[SPOT_COGS_Build_Tabs.xlsx]\" with nothing so links point at the model's own tabs.", italic=True)

    # ======================================================================= stubs
    for name, cells in [("RPM", {"B2": "STUB", **{f"{c}2": y for c, y in zip(RPM_COLS, YEARS)}, **{f"{c}58": v for c, v in zip(RPM_COLS, RPM_TOTAL)}, "B58": "Total Revenue"}),
                        ("Valuation", {"B2": "Case Selector", "B3": "Toggle", "C3": "Base", "E3": "STUB"}),
                        ("DM Inputs", {"B2": "STUB", **DMI}), ("IS", {"B2": "STUB", "AA3": "2025", "B12": "Cost of Revenue", "AA12": -11690})]:
        st = wb.create_sheet(name)
        for ref, v in cells.items(): st[ref] = v; st[ref].font = F_IN if isinstance(v, (int, float)) else F_I
    wb.save(OUT); print("saved", OUT)

if __name__ == "__main__":
    build()
